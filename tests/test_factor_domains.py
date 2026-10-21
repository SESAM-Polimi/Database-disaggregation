"""Projection properties and isolation of the copied machine input package."""
import math
from zipfile import ZipFile
import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from entice_inventory.reconciliation.factor_domains import project_factors, reconcile_prepared_factors


def test_projection_preserves_net_and_has_kkt_optimum():
    q = np.array([-.06, -.02, .025, .05])
    bounded = np.array([True, True, False, False])
    v = project_factors(q, bounded)
    assert np.all(v[:2] == 0)
    assert v.sum() == pytest.approx(q.sum(), abs=1e-16)
    assert (v[2:] < 0).all()  # Signed factors can absorb negative net VA.
    gradient = 2*(v-q)/abs(q)
    assert gradient[2] == pytest.approx(gradient[3])
    assert (gradient[:2] >= gradient[2]-1e-14).all()
    assert np.array_equal(q, [-.06, -.02, .025, .05])


def test_consistent_profile_is_untouched_and_zero_prior_is_explicit():
    q = [0., .2, -.1, .3]
    assert np.array_equal(project_factors(q, [True, True, False, False]), q)
    v = project_factors([-.1, 0., .3], [True, True, False])
    assert min(v[:2]) >= 0 and math.fsum(v) == pytest.approx(.2)
    with pytest.raises(ValueError, match='Negative net'):
        project_factors([-.1, .01], [True, True])
    with pytest.raises(ValueError, match='finite'):
        project_factors([float('nan')], [True])


def test_preparation_changes_only_selected_factor_cells(tmp_path):
    source = tmp_path/'source'; source.mkdir(); (source/'baseline').mkdir()
    (source/'baseline'/'immutable.bin').write_bytes(b'baseline')
    w = Workbook(); w.active.title = 'Master'
    w['Master'].append(['Region', 'Sector', 'Inventory sheet'])
    w['Master'].append(['BR', 'MCU', 'inventory']); w['Master'].append(['BR', 'OTHER', 'other'])
    for name in ['inventory', 'other']:
        s = w.create_sheet(name); s.append(['Quantity', 'Item type', 'DB Item'])
        s.append([.9, 'Sector', 'Parent'])
        for name_, q in [('Capital', -.06), ('Labor', -.02), ('TTM', .025), ('Tax', .05)]:
            s.append([q, 'Factor of production', name_])
    w.create_sheet('Total outputs').append([42])
    w.create_sheet('Intermediate steps').append(['MCU', 'OTHER'])
    w.save(source/'add_sectors.xlsx')
    original = (source/'add_sectors.xlsx').read_bytes()
    out = tmp_path/'reconciled'
    audit = reconcile_prepared_factors(source, out, ['MCU'])
    assert audit['profiles'][0]['region'] == 'BR'
    assert len(audit['changed_zip_members']) == 1
    assert (source/'add_sectors.xlsx').read_bytes() == original
    with ZipFile(source/'add_sectors.xlsx') as a, ZipFile(out/'add_sectors.xlsx') as b:
        for name in a.namelist():
            if name not in audit['changed_zip_members']:
                assert a.read(name) == b.read(name)
    result = load_workbook(out/'add_sectors.xlsx', data_only=True)
    assert result['inventory']['A2'].value == .9
    assert result['inventory']['A3'].value == 0
    assert result['other']['A3'].value == -.06
    result.close()
    with pytest.raises(ValueError, match='fresh output'):
        reconcile_prepared_factors(source, out, ['MCU'])
