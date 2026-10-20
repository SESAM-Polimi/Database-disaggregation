"""Geographic conservation, coefficient weights and shared-input preparation."""
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from entice_inventory.core.geography import Geography, aggregate_frame, aggregate_parquet
from entice_inventory.inventory.regionalize import aggregate_inventory


def inventory(path, *, sector='Child', outputs=(2., 6., 0.), profiles=None):
    book = Workbook()
    del book['Sheet']
    tables = {
        'Master': [['Region', 'Sector', 'Inventory sheet', 'Quantity', 'Unit', 'Final consumption',
                    'Consumption category', 'Parent Sector', 'Leave empty', 'Source', 'Notes', 'Add or Split']],
        'Regions Clusters': [['GLOBAL'], ['R1'], ['R2'], ['R3']],
        'Sectors Clusters': [['Cluster1']], 'Factors Clusters': [['Factor1']],
        'Total outputs': [['Sector', 'Region', 'Quantity', 'Unit', 'Source', 'Notes']],
        'Trades': [['Sector_from', 'Region_from', 'Region_to', 'Quantity', 'Unit', 'Source', 'Notes'],
                   [sector, 'R1', 'R2', 1., 'M USD'], [sector, 'R1', 'R3', 2., 'M USD'],
                   [sector, 'R2', 'R3', 3., 'M USD']],
    }
    for i, region in enumerate(('R1', 'R2', 'R3')):
        tables['Master'].append([region, sector, f'INV{i}', None, 'M USD', None, None, 'Parent', None, None, None, 'Split'])
        tables['Total outputs'].append([sector, region, outputs[i], 'M USD'])
        a = (.2, .6, .5)[i] if profiles is None else profiles[i]
        tables[f'INV{i}'] = [['Quantity', 'Unit', 'Input', 'Item type', 'DB Item', 'DB Region', 'Change type'],
                             [a, 'M USD', 'Parent input', 'Sector', 'Parent', 'GLOBAL', 'Update'],
                             [1-a, 'M USD', 'VA', 'Factor of production', 'VA', None, 'Update']]
    for name, rows in tables.items():
        ws = book.create_sheet(name)
        for row in rows:
            ws.append(row)
    book.save(path)
    return path


def test_mapping_requires_a_partition(tmp_path):
    path = tmp_path / 'mapping.csv'
    path.write_text('source,target\nR1,G\nR1,H\n')
    with pytest.raises(ValueError, match='Duplicate'):
        Geography.read(path)
    geo = Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'})
    assert geo.regions == ['G', 'H']
    with pytest.raises(ValueError, match='match the baseline'):
        geo.validate(['R1', 'R2'])
    json_path = tmp_path / 'mapping.json'
    json_path.write_text('{"R1":"G","R1":"H"}')
    with pytest.raises(ValueError, match='Duplicate'):
        Geography.read(json_path)


def test_camilla_partition_covers_the_canonical_gtap_universe():
    from entice_inventory.core.registry import Registry
    path = Path(__file__).resolve().parents[1] / 'data/geographies/gtap13.json'
    geo = Geography.read(path)
    geo.validate(Registry.load().members('GLOBAL', geography='GTAP12'))
    assert len(geo.regions) == 13


def test_chunked_baseline_matches_flow_aggregation(tmp_path):
    geo = Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'})
    axis = pd.MultiIndex.from_product([['R1', 'R2', 'R3'], ['Sector'], ['P', 'Q']], names=['Region', 'Level', 'Item'])
    frame = pd.DataFrame(np.arange(36, dtype=float).reshape(6, 6), index=axis, columns=axis)
    source = tmp_path / 'Z.parquet'
    frame.to_parquet(source)
    expected = aggregate_frame(frame, geo).sort_index().sort_index(axis=1)
    for size in (1, 3, 256):
        pd.testing.assert_frame_equal(aggregate_parquet(source, geo, chunk_size=size), expected)
    assert expected.to_numpy().sum() == frame.to_numpy().sum()


def test_output_weighted_profiles_and_internal_trade(tmp_path):
    source = inventory(tmp_path / 'native.xlsx')
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    target = tmp_path / 'aggregated.xlsx'
    geo = Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'})
    with pytest.warns(UserWarning, match='zero output'):
        report = aggregate_inventory(source, target, geo)
    book = load_workbook(target, data_only=True)
    assert report['output'] == 8.
    assert list(book['Total outputs'].values)[1][1:3] == ('G', 8.)
    assert book['INV_001']['A2'].value == pytest.approx(.5)
    assert list(book['Trades'].values)[1][1:4] == ('G', 'H', 5.)
    assert 'Exclusions' in book and 'Tolerances' in book
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    with pytest.raises(FileExistsError):
        aggregate_inventory(source, target, geo)


def test_unknown_output_is_not_a_zero_weight(tmp_path):
    source = inventory(tmp_path / 'native.xlsx', outputs=(None, 6., 0.))
    with pytest.raises(ValueError, match='Missing numeric'):
        aggregate_inventory(source, tmp_path / 'out.xlsx', Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'}))


def test_sparse_export_uses_explicit_native_observation(tmp_path):
    source = inventory(tmp_path / 'native.xlsx')
    book = load_workbook(source)
    book['Total outputs'].delete_rows(4)
    book.save(source)
    geo = Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'})
    with pytest.raises(ValueError, match='Missing native output'):
        aggregate_inventory(source, tmp_path / 'missing.xlsx', geo)
    reference = {'source': 'original inventory', 'rows': [dict(Sector='Child', Region='R3', Quantity=0., Unit='M USD')]}
    with pytest.warns(UserWarning, match='zero output'):
        report = aggregate_inventory(source, tmp_path / 'explicit.xlsx', geo, output_reference=reference)
    assert report['supplemented_output_regions'] == ['R3']
    assert report['output_reference'] == 'original inventory'


def test_partial_origin_cluster_is_not_expanded_silently(tmp_path):
    source = inventory(tmp_path / 'native.xlsx')
    book = load_workbook(source)
    book['Regions Clusters'].cell(1, 2, 'PARTIAL')
    book['Regions Clusters'].cell(2, 2, 'R1')
    book['INV0']['F2'] = 'PARTIAL'
    book.save(source)
    with pytest.raises(ValueError, match='partly covers'):
        aggregate_inventory(source, tmp_path / 'out.xlsx', Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'}))


def test_preparation_uses_same_geography_for_baseline_and_workbook(tmp_path):
    import mario
    from entice_inventory.core.disaggregation import prepare_inputs
    axis = pd.MultiIndex.from_product([['R1', 'R2', 'R3'], ['Sector'], ['Parent']], names=['Region', 'Level', 'Item'])
    demand = pd.MultiIndex.from_product([['R1', 'R2', 'R3'], ['Consumption category'], ['HH']], names=axis.names)
    z = pd.DataFrame(2., index=axis, columns=axis)
    y = pd.DataFrame(4., index=axis, columns=demand)
    v = pd.DataFrame([[12., 12., 12.]], index=['VA'], columns=axis)
    e = pd.DataFrame([[1., 2., 3.]], index=['CO2'], columns=axis)
    ey = pd.DataFrame(0., index=['CO2'], columns=demand)
    units = {kind: pd.DataFrame({'unit': [unit]}, index=[item]) for kind, item, unit in
             [('Sector', 'Parent', 'M USD'), ('Factor of production', 'VA', 'M USD'), ('Satellite account', 'CO2', 'tonne')]}
    db = mario.Database(name='Regional fixture', table='IOT', Z=z, Y=y, V=v, E=e, EY=ey, units=units)
    db.to_parquet(path=tmp_path / 'baseline', flows=True, coefficients=False)
    inputs = tmp_path / 'inputs'
    inputs.mkdir()
    inventory(inputs / 'Child.xlsx', outputs=(2., 6., 1.))
    geo = Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'})
    result = prepare_inputs(inputs, tmp_path / 'prepared', geography=geo,
                            steps={'Mining': ['Child']}, baseline_path=tmp_path / 'baseline' / 'flows')
    loaded = mario.parse_from_parquet(str(result / 'baseline' / 'flows'), table='IOT', mode='flows')
    loaded.read_add_sectors_excel(result / 'add_sectors.xlsx', read_inventories=True, split=True)
    assert set(loaded.get_index('Region')) == {'G', 'H'}
    assert set(loaded.add_sectors_master.Region) == {'G', 'H'}
    assert loaded.add_sectors_workbook.intermediate_steps == {'Mining': ['Child']}
    assert loaded.Z.to_numpy().sum() == db.Z.to_numpy().sum()
    assert json.loads((result / 'preparation.json').read_text())['geography']['mapping'] == geo.mapping
