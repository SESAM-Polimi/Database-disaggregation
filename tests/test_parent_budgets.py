"""Cumulative feasibility, directed routes, geography and blocking entry points."""
import json
from pathlib import Path
import runpy
import sys
from types import SimpleNamespace

import pandas as pd
import pytest

from entice_inventory.core.geography import Geography
from entice_inventory.core.parent_budgets import (
    ChildTargets, ParentBudgets, ParentBudgetError, evaluate_parent_budgets, require_parent_budgets,
)


@pytest.fixture
def budgets():
    index = pd.MultiIndex.from_product([['R1', 'R2', 'R3'], ['Parent']], names=['Region', 'Sector'])
    return ParentBudgets(pd.Series(100., index=index),
                         pd.DataFrame(10., index=index, columns=['R1', 'R2', 'R3']), {'Parent': 'M USD'})


@pytest.fixture
def targets():
    master = pd.DataFrame([['A', 'Parent', 'Split'], ['B', 'Parent', 'Split']],
                          columns=['Sector', 'Parent Sector', 'Add or Split'])
    output = pd.DataFrame([[s, r, q, 'M USD'] for s, q in [('A', 40.), ('B', 50.)]
                           for r in ['R1', 'R2', 'R3']], columns=['Sector', 'Region', 'Quantity', 'Unit'])
    trade = pd.DataFrame([['A', 'R1', 'R2', 4., 'M USD'], ['B', 'R1', 'R2', 5., 'M USD'],
                          ['A', 'R2', 'R1', 9., 'M USD']],
                         columns=['Sector_from', 'Region_from', 'Region_to', 'Quantity', 'Unit'])
    return ChildTargets(master, output, trade)


def test_reports_both_collective_excesses_even_when_each_child_fits(budgets, targets, tmp_path):
    targets.output.loc[targets.output.Region.eq('R1'), 'Quantity'] = 60.
    targets.trade.loc[targets.trade.Region_from.eq('R1'), 'Quantity'] = 6.
    with pytest.raises(ParentBudgetError) as caught:
        require_parent_budgets(budgets, targets, report_path=tmp_path/'checks.json')
    report = json.loads((tmp_path/'checks.json').read_text())
    assert caught.value.report == report
    assert report['summary']['output_violations'] == 1
    assert report['summary']['trade_violations'] == 1
    assert report['output_violations'][0]['child_sum'] == 120.
    assert report['trade_violations'][0]['origin'] == 'R1'
    assert report['trade_violations'][0]['destination'] == 'R2'
    assert report['trade_violations'][0]['child_sum'] == 12.


def test_equality_and_only_numerical_roundoff_are_accepted(budgets, targets):
    targets.output['Quantity'] = 50.
    targets.trade.loc[targets.trade.Region_from.eq('R1'), 'Quantity'] = 5.
    assert require_parent_budgets(budgets, targets)['status'] == 'passed'
    targets.output.loc[0, 'Quantity'] += 1e-7
    assert require_parent_budgets(budgets, targets)['status'] == 'passed'
    targets.output.loc[0, 'Quantity'] += 1e-4
    with pytest.raises(ParentBudgetError):
        require_parent_budgets(budgets, targets)


def test_missing_output_blocks_instead_of_becoming_zero(budgets, targets):
    partial = ChildTargets(targets.master, targets.output.iloc[1:], targets.trade)
    report = evaluate_parent_budgets(budgets, partial)
    assert report['summary']['output_violations'] == 0
    assert report['summary']['missing_output_cells'] == 1
    assert report['incomplete_outputs'][0]['missing_children'] == ['A']
    with pytest.raises(ParentBudgetError):
        require_parent_budgets(budgets, partial)


def test_regional_aggregation_does_not_hide_an_unknown_native_output(budgets, targets):
    partial = ChildTargets(targets.master, targets.output.iloc[1:], targets.trade)
    geo = Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'})
    report = evaluate_parent_budgets(budgets.aggregate(geo), partial.aggregate(geo))
    assert report['status'] == 'failed'
    assert report['summary']['missing_output_cells'] == 1
    assert report['incomplete_outputs'][0]['missing_native_regions'] == {'A': ['R1']}


def test_checks_the_selected_geography_and_excludes_internalized_trade(budgets, targets):
    targets.output.loc[targets.output.Region.eq('R1'), 'Quantity'] = 60.
    targets.output.loc[targets.output.Region.eq('R2'), 'Quantity'] = 5.
    targets.trade.loc[targets.trade.Region_from.eq('R1'), 'Quantity'] = 6.
    assert evaluate_parent_budgets(budgets, targets)['summary']['output_violations'] == 1
    geo = Geography({'R1': 'G', 'R2': 'G', 'R3': 'H'})
    report = require_parent_budgets(budgets.aggregate(geo), targets.aggregate(geo))
    assert report['summary']['regions'] == 2
    assert report['summary']['observed_trade_groups'] == 0
    external = pd.DataFrame([['A', 'R1', 'R3', 11., 'M USD'], ['B', 'R1', 'R3', 11., 'M USD']],
                             columns=targets.trade.columns)
    targets = ChildTargets(targets.master, targets.output, external)
    report = evaluate_parent_budgets(budgets.aggregate(geo), targets.aggregate(geo))
    assert report['summary']['trade_violations'] == 1
    assert report['trade_violations'][0]['parent_value'] == 20.


@pytest.mark.parametrize('problem', ['units', 'duplicate', 'nonfinite', 'unknown_sector'])
def test_invalid_target_coordinates_and_values_are_not_comparisons(budgets, targets, problem):
    if problem == 'units':
        targets.output.loc[0, 'Unit'] = 'USD'
    elif problem == 'duplicate':
        targets = ChildTargets(targets.master, pd.concat([targets.output, targets.output.iloc[:1]]), targets.trade)
    elif problem == 'nonfinite':
        targets.trade.loc[0, 'Quantity'] = float('nan')
    else:
        targets.trade.loc[0, 'Sector_from'] = 'Unknown'
    with pytest.raises(ValueError):
        require_parent_budgets(budgets, targets)


def test_resume_counts_only_pending_children(budgets, targets):
    current = SimpleNamespace(get_index=lambda _: ['Parent', 'A'])
    pending = targets.pending(current)
    assert set(pending.master.Sector) == {'B'}
    assert set(pending.output.Sector) == {'B'}
    assert set(pending.trade.Sector_from) == {'B'}
    assert require_parent_budgets(budgets, pending)['status'] == 'passed'


def _baseline(tmp_path):
    import mario
    axis = pd.MultiIndex.from_product([['R1', 'R2', 'R3'], ['Sector'], ['Parent']],
                                      names=['Region', 'Level', 'Item'])
    demand = pd.MultiIndex.from_product([['R1', 'R2', 'R3'], ['Consumption category'], ['HH']], names=axis.names)
    units = {kind: pd.DataFrame({'unit': [unit]}, index=[item]) for kind, item, unit in
             [('Sector', 'Parent', 'M USD'), ('Factor of production', 'VA', 'M USD'),
              ('Satellite account', 'CO2', 'tonne')]}
    db = mario.Database(name='Budget fixture', table='IOT', units=units,
                        Z=pd.DataFrame(2., index=axis, columns=axis), Y=pd.DataFrame(4., index=axis, columns=demand),
                        V=pd.DataFrame(12., index=['VA'], columns=axis), E=pd.DataFrame(1., index=['CO2'], columns=axis),
                        EY=pd.DataFrame(0., index=['CO2'], columns=demand))
    db.to_parquet(path=tmp_path/'baseline', flows=True, coefficients=False)
    return tmp_path/'baseline/flows'


def test_preparation_blocks_before_profiles_or_solver(tmp_path, monkeypatch):
    from test_geography import inventory
    from openpyxl import load_workbook
    from entice_inventory.core import disaggregation
    baseline = _baseline(tmp_path)
    source = tmp_path/'inputs'
    source.mkdir()
    for sector in ('A', 'B'):
        path = inventory(source/f'{sector}.xlsx', sector=sector, outputs=(14., 2., 1.))
        book = load_workbook(path)
        book['Trades']['D2'] = 4.
        book.save(path)
    def forbidden(*args, **kwargs):
        pytest.fail('Profile preparation must not start when cumulative targets exceed the parent.')
    monkeypatch.setattr(disaggregation, 'aggregate_inventory', forbidden)
    destination = tmp_path/'prepared'
    with pytest.raises(ParentBudgetError) as caught:
        disaggregation.prepare_inputs(source, destination, baseline_path=baseline, intermediate_steps=True)
    assert not destination.exists()
    assert caught.value.report['summary']['output_violations'] == 1
    assert caught.value.report['summary']['trade_violations'] == 1
    assert list(tmp_path.glob('.prepared_*/parent_budgets.json'))


def test_edited_prepared_workbook_is_rechecked_before_solve(tmp_path, monkeypatch):
    from test_geography import inventory
    import mario
    from openpyxl import load_workbook
    from entice_inventory.core.disaggregation import prepare_inputs
    baseline = _baseline(tmp_path)
    source = tmp_path/'inputs'
    source.mkdir()
    inventory(source/'A.xlsx', sector='A', outputs=(2., 6., 1.))
    prepared = prepare_inputs(source, tmp_path/'prepared', baseline_path=baseline, intermediate_steps=True)
    book = load_workbook(prepared/'add_sectors.xlsx')
    book['Total outputs']['C2'] = 30.
    book.save(prepared/'add_sectors.xlsx')
    def forbidden(*args, **kwargs):
        pytest.fail('The solver must not run after an edit exceeding the parent.')
    monkeypatch.setattr(mario.Database, 'add_sectors', forbidden)
    monkeypatch.setattr(sys, 'argv', ['run_disaggregation.py', '--prepared', str(prepared), '--solve'])
    script = Path(__file__).resolve().parents[1]/'scripts/run_disaggregation.py'
    with pytest.raises(ParentBudgetError):
        runpy.run_path(str(script), run_name='__main__')
    assert not (prepared/'checkpoints').exists()
    assert (prepared/'parent_budgets_solve.json').exists()
