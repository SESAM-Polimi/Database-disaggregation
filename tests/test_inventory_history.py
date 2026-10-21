"""Cumulative publication keeps earlier profiles and later reconciled targets."""
from dataclasses import replace
import hashlib
import json
from pathlib import Path

import pandas as pd
import pytest

from entice_inventory.core.inventory_history import (
    assemble_inventory_history, checkpoint_history,
)


def payload(sectors=('A', 'B'), *, trade=2.):
    from mario.ops.add_sector_workbook import AddSectorWorkbook, build_add_sector_master_sheet
    master = build_add_sector_master_sheet('IOT', list(sectors), ['R1', 'R2', 'R3'])
    master['Parent Sector'] = 'Parent'
    master['Unit'] = 'M USD'
    master['Add or Split'] = 'Split'
    inventories = {}
    for row in master.to_dict('records'):
        # A already contains a finer observation for its future supplier B.
        records = [[.2, 'M USD', 'input', 'Sector', 'Parent', 'GLOBAL', 'Update'],
                   [.8, 'M USD', 'VA', 'Factor of production', 'VA', None, 'Update']]
        if row['Sector'] == 'A':
            records.append([.1, 'M USD', 'future B', 'Sector', 'B', 'GLOBAL', 'Update'])
        inventories[row['Inventory sheet']] = pd.DataFrame(records, columns=[
            'Quantity', 'Unit', 'Input', 'Item type', 'DB Item', 'DB Region', 'Change type'])
    support = {
        'Total outputs': pd.DataFrame([[s, r, 0., 'M USD', None, None] for s in sectors
                                      for r in ['R1', 'R2', 'R3']],
                                     columns=['Sector', 'Region', 'Quantity', 'Unit', 'Source', 'Notes']),
        'Trades': pd.DataFrame([[s, 'R1', 'R2', trade, 'M USD', None, None] for s in sectors],
                              columns=['Sector_from', 'Region_from', 'Region_to', 'Quantity', 'Unit', 'Source', 'Notes']),
        'Exclusions': pd.DataFrame(columns=['Region_from', 'Region_to', 'Sector_from', 'Sector_to', 'Notes']),
        'Tolerances': pd.DataFrame({'tol_Name': ['delta', 'eps'], 'values': [0., 0.]}),
    }
    return AddSectorWorkbook('IOT', master, {'GLOBAL': ['R1', 'R2', 'R3']}, {}, {}, {},
                             inventories, support, {f'Step {s}': [s] for s in sectors})


def make_trial(root, source, *, active, completed, index, previous=None):
    """Synthetic accepted files; no optimisation is performed by this fixture."""
    import mario
    from test_parent_budgets import _baseline
    from mario.ops.add_sector_steps import compile_step, write_workbook_payload
    _baseline(root)
    base = mario.parse_from_parquet(str(root/'baseline/flows'), table='IOT', mode='flows')
    trial = root/'trial'; trial.mkdir()
    write_workbook_payload(source, trial/'add_sectors.xlsx')
    step = trial/f'checkpoints/{index:03d}_Test'; step.mkdir(parents=True)
    parents = {s: 'Parent' for s in set(completed) | set(source.master_sheet.Sector)}
    existing = {'Parent'} | (set(completed) - set(active))
    compiled = compile_step(source, active, existing, parents)
    write_workbook_payload(compiled, step/'step_add_sectors.xlsx')
    write_workbook_payload(source, step/'source_add_sectors.xlsx')
    pending = source.master_sheet.loc[~source.master_sheet.Sector.isin(completed)]
    remaining = replace(source, master_sheet=pending,
                        intermediate_steps={k: [s for s in v if s not in completed]
                                            for k, v in source.intermediate_steps.items() if any(s not in completed for s in v)})
    write_workbook_payload(remaining, step/'remaining_add_sectors.xlsx')
    axis = pd.MultiIndex.from_product([['R1', 'R2', 'R3'], ['Sector'], ['Parent', *completed]],
                                     names=base.Z.index.names)
    units = dict(base.units)
    units['Sector'] = pd.DataFrame({'unit': 'M USD'}, index=['Parent', *completed])
    db = mario.Database(name='Published fixture', table='IOT', units=units,
                        Z=base.Z.reindex(index=axis, columns=axis, fill_value=0.),
                        Y=base.Y.reindex(index=axis, fill_value=0.),
                        V=base.V.reindex(columns=axis, fill_value=0.),
                        E=base.E.reindex(columns=axis, fill_value=0.), EY=base.EY)
    db.to_parquet(path=step/'database', flows=True, coefficients=False)
    state = dict(format='mario-sector-steps-v1', status='complete', step_index=index,
                 step_name=f'Step {active[0]}', step_sectors=active, completed_sectors=completed,
                 sector_parents=parents,
                 files={str(p.relative_to(step)): hashlib.sha256(p.read_bytes()).hexdigest()
                        for p in (step/'database').rglob('*') if p.is_file()})
    (step/'state.json').write_text(json.dumps(state))
    manifest = dict(status='accepted', started_at=f'2026-10-21T04:{54+index}:00+00:00',
                    source_preparation=str(root), accepted_checkpoints=[str(step)])
    if previous:
        manifest['source_checkpoint'] = str(previous)
    (trial/'trial.json').write_text(json.dumps(manifest))
    return trial, step


def test_resumed_publication_is_complete_and_preserves_exact_stage_inputs(tmp_path, monkeypatch):
    import mario
    from mario.ops.add_sector_workbook import load_add_sector_workbook_input
    from entice_inventory.core.shared_runs import publish_trial
    monkeypatch.setattr(mario.Database, 'add_sectors', lambda *a, **k: pytest.fail('Publication must not solve.'))
    first_root = tmp_path/'first'; first_root.mkdir()
    first, first_step = make_trial(first_root, payload(), active=['A'], completed=['A'], index=1)
    options = dict(inventory_root=tmp_path/'Inventories', database_root=tmp_path/'Database')
    first_result = publish_trial(first, **options)
    second_root = tmp_path/'second'; second_root.mkdir()
    second, second_step = make_trial(second_root, payload(('B',), trade=7.),
                                    active=['B'], completed=['A', 'B'], index=2,
                                    previous=first_result['database'])
    original = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in tmp_path.rglob('*') if p.is_file()}
    result = publish_trial(second, **options)
    inv = Path(result['inventory'])
    merged, _ = load_add_sector_workbook_input(inv/'add_sectors.xlsx', table='IOT', read_inventories=True, split=True)
    assert set(merged.master_sheet.Sector) == {'A', 'B'}
    assert merged.intermediate_steps == {'Step A': ['A'], 'Step B': ['B']}
    assert merged.split_info['Trades'].set_index('Sector_from').Quantity.to_dict() == {'A': 2., 'B': 7.}
    assert len(merged.master_sheet) == len(merged.inventories_by_sheet) == 6
    assert (inv/'step_add_sectors.xlsx').read_bytes() == (second_step/'step_add_sectors.xlsx').read_bytes()
    assert (inv/'source_add_sectors.xlsx').read_bytes() == (second/'add_sectors.xlsx').read_bytes()
    assert json.loads((Path(result['database'])/'run.json').read_text())['source_database_directory'].startswith('../')
    assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest() == h for p, h in original.items())
    # A partial publication exposes the full editable plan, including pending B.
    first_manifest = json.loads((Path(first_result['inventory'])/'inventory_manifest.json').read_text())
    assert first_manifest['completed_sectors'] == ['A']
    assert set(first_manifest['planned_sectors']) == {'A', 'B'}


def test_missing_prior_history_fails_instead_of_publishing_partial_cumulative_inputs(tmp_path):
    _, step = make_trial(tmp_path, payload(('B',)), active=['B'], completed=['A', 'B'], index=2)
    destination = tmp_path/'bad.xlsx'
    with pytest.raises(ValueError, match='Incomplete'):
        assemble_inventory_history(checkpoint_history(step), destination)
    assert not destination.exists()


def test_changed_saved_step_is_detected_before_publication(tmp_path):
    from openpyxl import load_workbook
    _, step = make_trial(tmp_path, payload(), active=['A'], completed=['A'], index=1)
    book = load_workbook(step/'step_add_sectors.xlsx')
    book['Trades']['D2'] = 99.
    book.save(step/'step_add_sectors.xlsx')
    with pytest.raises(AssertionError):
        assemble_inventory_history(checkpoint_history(step), tmp_path/'bad.xlsx')


def test_upgrade_archives_legacy_input_and_does_not_touch_database(tmp_path):
    import runpy
    from entice_inventory.core.shared_runs import publish_trial
    trial, step = make_trial(tmp_path, payload(('New',)), active=['New'], completed=['New'], index=1)
    result = publish_trial(trial, inventory_root=tmp_path/'Inventories', database_root=tmp_path/'Database')
    inventory, database = Path(result['inventory']), Path(result['database'])
    # Recreate the former publication layout, with the compiled step as primary.
    (inventory/'add_sectors.xlsx').write_bytes((inventory/'step_add_sectors.xlsx').read_bytes())
    (inventory/'step_add_sectors.xlsx').unlink()
    (inventory/'inventory_manifest.json').unlink()
    previous = (inventory/'add_sectors.xlsx').read_bytes()
    before = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in database.rglob('*') if p.is_file()}
    upgrade = runpy.run_path(str(Path(__file__).resolve().parents[1]/'scripts/upgrade_published_inventories.py'))['upgrade']
    upgrade(database)
    assert (inventory/'step_add_sectors.xlsx').read_bytes() == previous
    assert pd.read_excel(inventory/'add_sectors.xlsx', sheet_name='Intermediate steps').iloc[0, 0] == 'New'
    assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest() == h for p, h in before.items())
    with pytest.raises(FileExistsError):
        upgrade(database)
