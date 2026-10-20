"""Economic assumptions, independent capacity checks and replayable preparation."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook,load_workbook

from entice_inventory.core.geography import Geography
from entice_inventory.core.parent_budgets import ParentBudgets
from entice_inventory.reconciliation.audit import Audit
from entice_inventory.reconciliation.baseline import NativeBaseline
from entice_inventory.reconciliation.policy import load_policy
from entice_inventory.reconciliation.profiles import balance_profile
from entice_inventory.reconciliation.sources import Sources
from entice_inventory.reconciliation.targets import prepare_outputs,reconcile_trade


def toy_policy(codes=('A','B')):
    policy,_=load_policy()
    policy['id']='toy-cleaning-v1';policy['require_all_sectors']=True
    policy['sectors']={c:{'name':c,'declared_parent':'Parent',
        'output_source':{'kind':'legacy_share','assumption':'Synthetic legacy share'}} for c in codes}
    policy['overlap_groups']=[];policy['trade']['composites']={};policy['trade']['parent_pattern_only']=[]
    return policy


def budgets(regions=('R1','R2','R3'),output=10.,trade=5.):
    axis=pd.MultiIndex.from_product([regions,['Parent']],names=['Region','Sector'])
    return ParentBudgets(pd.Series(output,index=axis),pd.DataFrame(trade,index=axis,columns=regions),{'Parent':'M USD'})


def test_policy_freezes_all_67_explicit_decisions_and_rejects_unsupported_rules():
    policy,digest=load_policy()
    assert len(policy['sectors'])==67 and len(digest)==64
    assert policy['sectors']['MSI']['output_source']['factor']==1e-6
    assert policy['sectors']['PLSR']['output_source']['column']=='E'
    assert policy['sectors']['RBR']['output_source']['kind']=='template_share'
    broken=deepcopy(policy);broken['overlap_groups'][0]['weights']=[1.,1.]
    with pytest.raises(ValueError,match='sum to one'):load_policy(broken)
    broken=deepcopy(policy);broken['profiles']['negative_factors']='zero'
    with pytest.raises(ValueError,match='Signed'):load_policy(broken)


def test_native_streaming_budgets_and_parent_profile(tmp_path):
    from test_parent_budgets import _baseline
    import mario
    path=_baseline(tmp_path)
    native=NativeBaseline(path,chunk_size=1)
    db=mario.parse_from_parquet(str(path),table='IOT',mode='flows')
    expected=ParentBudgets.from_database(db)
    pd.testing.assert_series_equal(native.budgets.output.sort_index(),expected.output.sort_index(),check_names=False)
    pd.testing.assert_frame_equal(native.budgets.trade.sort_index(),expected.trade.sort_index())
    profile,donors=native.profile(('R1','R2'),'Parent')
    assert donors==('R1','R2')
    assert sum(r['Quantity'] for r in profile if r['Item type']!='Satellite account')==pytest.approx(1.)


def test_weighted_monetary_balance_preserves_signed_tax_and_satellites(tmp_path):
    audit=Audit(tmp_path/'audit');policy=toy_policy()
    profile=[{'Quantity':q,'Unit':unit,'Input':item,'DB Item':item,'Item type':kind,'DB Region':origin,'Change type':'Update'}
             for item,kind,q,origin,unit in [('Parent','Sector',.8,'GLOBAL','M USD'),
                 ('Labor','Factor of production',.4,None,'M USD'),('Tax','Factor of production',-.1,None,'M USD'),
                 ('CO2','Satellite account',42.,None,'tonne')]]
    result=balance_profile(profile,policy,audit,sector='A',region='R1');audit.close()
    values={r['DB Item']:r['Quantity'] for r in result}
    assert values['Tax']==-.1 and values['CO2']==42.
    assert sum(values[i] for i in ['Parent','Labor','Tax'])==pytest.approx(1.)
    assert abs(values['Labor']-.4)>abs(values['Parent']-.8)
    assert profile[0]['Quantity']==.8


def test_trade_fit_respects_producer_and_route_budgets_together(tmp_path):
    b=budgets();policy=toy_policy();audit=Audit(tmp_path/'audit')
    priors={('A','R1','R2'):4.,('A','R1','R3'):4.,('B','R1','R2'):4.,('B','R1','R3'):.2}
    kinds={k:'direct' for k in priors}
    result,status=reconcile_trade(priors,kinds,{'A':'Parent','B':'Parent'},
                                  {('A','R1'):3.,('B','R1'):3.},b,policy,audit)
    audit.close()
    assert sum(q for (c,o,d),q in result.items() if c=='A')<=3.+1e-8
    assert sum(q for (c,o,d),q in result.items() if c=='B')<=3.+1e-8
    assert sum(q for (c,o,d),q in result.items() if d=='R2')<=5.+1e-8
    assert all(0<=q<=priors[k]+1e-9 for k,q in result.items())
    assert status[0]['status'] in ('optimal','optimal_inaccurate')


def test_reliable_trade_is_changed_less_and_zero_capacity_is_structural(tmp_path):
    b=budgets(trade=10.);policy=toy_policy();audit=Audit(tmp_path/'audit')
    priors={('A','R1','R2'):10.,('B','R1','R2'):10.,('A','R2','R3'):4.}
    kinds={('A','R1','R2'):'paired',('B','R1','R2'):'imputed',('A','R2','R3'):'direct'}
    result,_=reconcile_trade(priors,kinds,{'A':'Parent','B':'Parent'},
                             {('A','R1'):20.,('B','R1'):20.,('A','R2'):0.},b,policy,audit)
    audit.close()
    assert result[('A','R1','R2')]>result[('B','R1','R2')]
    assert result[('A','R2','R3')]==0.


def test_missing_native_output_uses_parent_share_and_overlap_total_is_used_once(tmp_path):
    policy=toy_policy();b=budgets(regions=('R1','R2'))
    native=SimpleNamespace(regions=['R1','R2'],budgets=b)
    geo=Geography({'R1':'G','R2':'G'})
    observations=lambda c,p:[{'members':['R1'],'value':2. if c=='A' else 4.,'identifier':'R1',
                              'source':'synthetic','state':'numeric','kind':'direct','confidence':1.}]
    source=SimpleNamespace(observations=observations)
    from entice_inventory.core.parent_budgets import ChildTargets
    target=ChildTargets(pd.DataFrame([['A'],['B']],columns=['Sector']),
                        pd.DataFrame(columns=['Sector','Region','Quantity','Unit']),pd.DataFrame())
    audit=Audit(tmp_path/'audit')
    native_values,final,_=prepare_outputs(policy,source,native,geo,target,{'A':'Parent','B':'Parent'},
                                         {'A':'Parent','B':'Parent'},audit)
    assert native_values[('A','R2')]==2. and final[('A','G')]==4.
    policy['overlap_groups']=[{'id':'one_observation','reference':'B','members':['A','B'],
                              'weights':[.5,.5],'assumption':'Synthetic one-time allocation'}]
    native_values,final,_=prepare_outputs(policy,source,native,geo,target,{'A':'Parent','B':'Parent'},
                                         {'A':'Parent','B':'Parent'},audit)
    audit.close()
    assert final[('A','G')]+final[('B','G')]==8.


def _registry():
    return SimpleNamespace(regions=lambda:[{'id':'R1','name':'One'},{'id':'R2','name':'Two'}],members=lambda v:[])


def test_template_column_scale_blank_and_zero_are_distinct(tmp_path):
    policy=toy_policy(('A',));policy['sectors']['A']['output_source']={
        'kind':'template','file':'A.xlsx','column':'G','unit_column':'H','factor':1e-6,'assumption':'literal USD'}
    book=Workbook();del book['Sheet'];ws=book.create_sheet('Region')
    ws.append(['Group','Name','Code','present']);ws.append([None]*4)
    ws.append(['G','One','R1',1]);ws.append(['G','Two','R2',1])
    ws=book.create_sheet('Total production');ws.append(['Region',None,None,None,'2023 Value',None,'2023 Value','Unit'])
    ws.append([None]*8);ws.append(['One',None,None,None,999.,None,2e6,'USD'])
    ws.append(['Two',None,None,None,999.,None,None,'USD'])
    book.save(tmp_path/'A.xlsx');book.close()
    audit=Audit(tmp_path/'audit');native=SimpleNamespace(regions=['R1','R2'])
    sources=Sources({'partner_templates':tmp_path,'purdue_trade':None},_registry(),native,policy,audit)
    data=sources.observations('A','Parent');audit.close()
    assert len(data)==1 and data[0]['value']==2.
    events=[json.loads(v) for v in (tmp_path/'audit/changes.jsonl').read_text().splitlines()]
    assert any(r.get('identifier')=='Two' and r.get('state')=='blank' for r in events)


def test_sparse_purdue_zeros_require_declared_domain(tmp_path):
    policy=toy_policy(('A',));policy['sectors']['A']['output_source']={'kind':'purdue','assumption':'paired test'}
    book=Workbook();del book['Sheet']
    for name,rows in {'REG':[['REG'],['R1'],['R2']], 'REPSEC':[['REPSEC'],['A']],
                      'REPOUT':[['r','s','value'],['R1','A',2.]],'BLTTRD':[['o','d','s','value']]}.items():
        ws=book.create_sheet(name)
        for row in rows:ws.append(row)
    book.save(tmp_path/'paired.xlsx');book.close()
    audit=Audit(tmp_path/'audit');native=SimpleNamespace(regions=['R1','R2'])
    sources=Sources({'purdue_split_targets':tmp_path/'paired.xlsx','purdue_trade':None},_registry(),native,policy,audit)
    data=sources.observations('A','Parent')
    assert data[1]['value']==0 and data[1]['state']=='implicit_parameter_zero'
    policy['sectors']['B']={'output_source':{'kind':'purdue','assumption':'not in domain'}}
    with pytest.raises(ValueError,match='absent'):sources.observations('B','Parent')
    audit.close()


def test_reconciled_preparation_is_readable_auditable_and_preserves_sources(tmp_path):
    from test_parent_budgets import _baseline
    from test_geography import inventory
    from entice_inventory.core.disaggregation import prepare_inputs
    from entice_inventory.core.parent_budgets import read_child_targets,require_parent_budgets
    import mario
    baseline=_baseline(tmp_path);source=tmp_path/'source';source.mkdir()
    for sector in ['A','B']:inventory(source/f'{sector}.xlsx',sector=sector,outputs=(40.,40.,40.))
    hashes={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in list(source.glob('*.xlsx'))+list(baseline.glob('*.parquet'))}
    prepared=prepare_inputs(source,tmp_path/'prepared',geography=Geography({'R1':'G','R2':'G','R3':'H'}),
        steps={'First':['A'],'Then':['B']},cleaning_policy=toy_policy(),baseline_path=baseline,
        source_paths={'purdue_trade':None,'purdue_split_targets':None,'partner_templates':None})
    manifest=json.loads((prepared/'preparation.json').read_text())
    assert manifest['reconciliation']['status']=='complete'
    assert manifest['reconciliation']['source_certified'] is False
    assert manifest['reconciliation']['runtime']['cvxlab']=='1.1.0'
    assert manifest['reconciliation']['checks_before']['output_violations']>0
    assert manifest['reconciliation']['checks_after']['output_violations']==0
    assert manifest['reconciliation']['checks_after']['producer_export_violations']==0
    assert (prepared/'reconciled/README.md').exists()
    assert (prepared/'reconciled/audit/policy.json').exists()
    assert (prepared/'reconciled/audit/changes.jsonl').stat().st_size>0
    db=mario.parse_from_parquet(str(prepared/'baseline/flows'),table='IOT',mode='flows')
    db.read_add_sectors_excel(prepared/'add_sectors.xlsx',read_inventories=True,split=True)
    assert db.add_sectors_workbook.intermediate_steps=={'First':['A'],'Then':['B']}
    require_parent_budgets(ParentBudgets.from_database(db),read_child_targets(sorted((prepared/'reconciled/inventories').glob('*.xlsx'))))
    assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==v for p,v in hashes.items())
    with pytest.raises(FileExistsError):
        prepare_inputs(source,prepared,cleaning_policy=toy_policy(),baseline_path=baseline)


def test_share_proxy_is_a_weak_parent_share_and_changed_headers_stop(tmp_path):
    policy=toy_policy(('A',));policy['sectors']['A']['output_source']={
        'kind':'template_share','file':'A.xlsx','column':'B','expected_headers':['Share',None],
        'assumption':'Residual-country share proxy'}
    policy['geography']['region_proxies']={'A':{'Small country':{'target':'R2','confidence':.1}}}
    book=Workbook();del book['Sheet'];book.create_sheet('Region')
    ws=book.create_sheet('Total production');ws.append(['Region','Share']);ws.append([None,None])
    ws.append(['Small country',.25]);book.save(tmp_path/'A.xlsx');book.close()
    audit=Audit(tmp_path/'audit');native=SimpleNamespace(regions=['R1','R2'],budgets=budgets(('R1','R2')))
    sources=Sources({'partner_templates':tmp_path,'purdue_trade':None},_registry(),native,policy,audit)
    observation=sources.observations('A','Parent')[0]
    assert observation['members']==['R2'] and observation['value']==2.5 and observation['confidence']==.1
    book=load_workbook(tmp_path/'A.xlsx')
    assert sources._clusters(book)['WORLD']=={'R1','R2'}
    book.close()
    audit.close()
    policy['sectors']['A']['output_source']['expected_headers']=['Different measure',None]
    audit=Audit(tmp_path/'other_audit');sources.audit=audit
    with pytest.raises(ValueError,match='headers changed'):sources.observations('A','Parent')
    audit.close()


def test_residual_world_total_adds_to_direct_country_and_parent_zero_stays_zero(tmp_path):
    from entice_inventory.core.parent_budgets import ChildTargets
    policy=toy_policy(('A',));policy['geography']['cluster_modes']={'A':{'GLOBAL':{'mode':'residual'}}}
    b=budgets(('R1','R2','R3'));b.output.loc[('R3','Parent')]=0.
    native=SimpleNamespace(regions=['R1','R2','R3'],budgets=b)
    observations=[{'members':['R1'],'value':2.,'identifier':'R1','source':'test',
                   'state':'numeric','kind':'direct','confidence':1.},
                  {'members':native.regions,'value':3.,'identifier':'GLOBAL','source':'test',
                   'state':'numeric','kind':'cluster','confidence':1.}]
    target=ChildTargets(pd.DataFrame([['A']],columns=['Sector']),
                        pd.DataFrame(columns=['Sector','Region','Quantity','Unit']),pd.DataFrame())
    audit=Audit(tmp_path/'audit')
    values,final,_=prepare_outputs(policy,SimpleNamespace(observations=lambda *_:observations),native,
        Geography({r:r for r in native.regions}),target,{'A':'Parent'},{'A':'Parent'},audit)
    audit.close()
    assert final[('A','R1')]==pytest.approx(2.,abs=.01)
    assert final[('A','R2')]==pytest.approx(3.,abs=.01)
    assert final[('A','R3')]==0.
    events=[json.loads(v) for v in (tmp_path/'audit/changes.jsonl').read_text().splitlines()]
    residual=next(e for e in events if e['rule']=='joint_geographic_fit' and e['identifier']=='GLOBAL')
    assert residual['members']==['R2','R3']


def test_composite_trade_is_allocated_once_after_direct_members(tmp_path):
    from entice_inventory.reconciliation.targets import prepare_trade
    policy=toy_policy(('A','B','C'));policy['trade']['composites']={'Composite':['A','B','C']}
    native=SimpleNamespace(regions=['R1','R2'],budgets=budgets(('R1','R2')))
    sources=SimpleNamespace(domain_sectors=set(),paired_trade={},customs_trade={
        'A':{('R1','R2'):2.},'Composite':{('R1','R2'):10.}})
    values={(c,r):q for c,q in [('A',2.),('B',1.),('C',3.)] for r in native.regions}
    audit=Audit(tmp_path/'audit')
    priors,kinds=prepare_trade(policy,sources,native,Geography({r:r for r in native.regions}),
                               {c:'Parent' for c in ['A','B','C']},values,values,audit)
    audit.close()
    assert priors=={('A','R1','R2'):2.,('B','R1','R2'):2.,('C','R1','R2'):6.}
    assert kinds[('B','R1','R2')]=='composite'


def test_missing_technology_uses_same_child_donor_and_reparent_rebuilds(tmp_path):
    from test_geography import inventory
    from test_parent_budgets import _baseline
    from entice_inventory.reconciliation.profiles import ProfileBuilder
    baseline=_baseline(tmp_path);native=NativeBaseline(baseline)
    source=tmp_path/'A.xlsx';inventory(source,sector='A',outputs=(1.,1.,1.))
    book=load_workbook(source);master=book['Master']
    master.delete_rows(3,master.max_row-2);book.save(source);book.close()
    geo=Geography({'R1':'G','R2':'G','R3':'H'});policy=toy_policy(('A',))
    audit=Audit(tmp_path/'audit')
    builder=ProfileBuilder(native,geo,{'A':'Parent'},{('A',r):1. for r in native.regions},policy,audit)
    profiles,_=builder.build(source,'A','Parent')
    assert profiles['G']==profiles['H']
    rebuilt,_=builder.build(source,'A','Former parent')
    audit.close()
    assert sum(r['Quantity'] for r in rebuilt['G'] if r['Item type']!='Satellite account')==pytest.approx(1.)
    events=[json.loads(v) for v in (tmp_path/'audit/changes.jsonl').read_text().splitlines()]
    assert any(e['rule']=='profile_donor' and e['mode']=='same_child_global' for e in events)
    assert len([e for e in events if e['rule']=='reparented_technology'])==2


def test_failed_run_keeps_evidence_without_success_marker(tmp_path):
    from test_parent_budgets import _baseline
    from test_geography import inventory
    from entice_inventory.reconciliation import reconcile_inventories
    source=tmp_path/'source';source.mkdir();inventory(source/'A.xlsx',sector='A')
    original=hashlib.sha256((source/'A.xlsx').read_bytes()).hexdigest()
    policy=toy_policy(('A',));policy['sectors']['A']['declared_parent']='Different parent'
    with pytest.raises(ValueError,match='parent changed'):
        reconcile_inventories(source,tmp_path/'failed',policy=policy,baseline_path=_baseline(tmp_path),
                             source_paths={'purdue_trade':None})
    assert not (tmp_path/'failed/complete.json').exists()
    assert json.loads((tmp_path/'failed/audit/summary.json').read_text())['status']=='failed'
    assert (tmp_path/'failed/failure.json').exists()
    assert hashlib.sha256((source/'A.xlsx').read_bytes()).hexdigest()==original


def test_changing_source_during_audit_stops_replay(tmp_path):
    source=tmp_path/'source.txt';source.write_text('original')
    audit=Audit(tmp_path/'audit');audit.read(source,role='test')
    source.write_text('changed')
    with pytest.raises(ValueError,match='changed during'):audit.read(source,role='test')
    audit.close()
