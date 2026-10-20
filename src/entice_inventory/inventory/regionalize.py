"""Aggregate already resolved regional cost profiles, with explicit output weights.

This is the boundary after native inventory resolution / D2.4 export. Missing
rows in these resolved, sparse profiles represent zero coefficients. Unresolved
partial overrides must be resolved upstream before calling this adapter.
"""
from __future__ import annotations

from collections import defaultdict
import hashlib
from io import BytesIO
import math
from pathlib import Path
import tempfile
import warnings

from openpyxl import load_workbook


def _rows(ws):
    header, *rows = ws.values
    return [dict(zip(header, row)) for row in rows if any(v is not None for v in row)]


def _write(book, name, rows, columns):
    if name in book:
        del book[name]
    ws = book.create_sheet(name)
    ws.append(columns)
    for row in rows:
        ws.append([row.get(c) for c in columns])


def _number(value, context):
    if value is None or isinstance(value, bool):
        raise ValueError(f"Missing numeric value: {context}")
    value = float(value)
    if not math.isfinite(value):
        raise ValueError(f"Non-finite value: {context}")
    return value


def aggregate_inventory(source, target, geography, *, output_reference=None):
    """Write one resolved workbook in the requested geography; never edit source."""
    source, target = Path(source), Path(target)
    if source.resolve() == target.resolve() or target.exists():
        raise FileExistsError(f"Refusing to overwrite an inventory: {target}")
    # Read sequentially before ZIP seeking: cloud placeholders may not yet
    # expose a usable central directory on the first random-access read.
    raw = source.read_bytes()
    book = load_workbook(BytesIO(raw), data_only=True)
    try:
        master_columns = [v.value for v in book["Master"][1]]
        master = _rows(book["Master"])
        sectors = {r["Sector"] for r in master}
        if len(sectors) != 1:
            raise ValueError("Regional aggregation expects one sector per resolved inventory workbook.")
        sector = sectors.pop()
        clusters = {col[0]: [v for v in col[1:] if v is not None]
                    for col in zip(*book["Regions Clusters"].values) if col[0] is not None}
        outputs, output_units = {}, set()
        for row in _rows(book["Total outputs"]):
            if row["Sector"] != sector:
                raise ValueError("Unexpected sector in Total outputs.")
            region = row["Region"]
            if region in outputs or region not in geography.mapping:
                raise ValueError(f"Duplicate or unmapped output region: {region}")
            outputs[region] = _number(row["Quantity"], f"{sector}/{region} output")
            if outputs[region] < 0:
                raise ValueError("Output weights cannot be negative.")
            output_units.add(row["Unit"])
        supplemented = []
        if output_reference is not None:
            for row in output_reference['rows']:
                region = row['Region']
                if row['Sector'] == sector and region not in outputs and region in geography.mapping:
                    value = _number(row['Quantity'], f"{sector}/{region} reference output")
                    if value < 0:
                        raise ValueError('Reference output weights cannot be negative.')
                    outputs[region] = value
                    output_units.add(row['Unit'])
                    supplemented.append(region)
        missing = geography.mapping.keys() - outputs.keys()
        if missing:
            raise ValueError(f"Missing native output observations for {sector}: {sorted(missing)}; "
                             "missing production is not inferred as zero.")
        if len(output_units) != 1:
            raise ValueError("Output units must agree before regional aggregation.")
        profiles, settings = {}, {}
        for row in master:
            region = row["Region"]
            members = [region] if region in geography.mapping else clusters.get(region, [])
            if not members or set(members) - geography.mapping.keys():
                raise ValueError(f"Unresolved producing-region cluster: {region}")
            for member in members:
                if member in profiles:
                    raise ValueError(f"Overlapping inventory profiles for {sector}/{member}")
                profiles[member] = _rows(book[row["Inventory sheet"]])
                settings[member] = row
        missing_profiles = [r for r, x in outputs.items() if x > 0 and not profiles.get(r)]
        if missing_profiles:
            raise ValueError(f"Positive outputs lack resolved profiles: {missing_profiles}")

        # Cluster updates must denote unions of complete target regions.
        mapped_clusters = {}
        for name, members in clusters.items():
            if set(members) - geography.mapping.keys():
                raise ValueError(f"Unknown members of origin cluster {name}")
            targets = list(dict.fromkeys(geography.mapping[r] for r in members))
            covered = {r for r, g in geography.mapping.items() if g in targets}
            mapped_clusters[name] = targets if covered == set(members) else None

        new_master, new_outputs, inventory_sheets = [], [], []
        for group in geography.regions:
            members = [r for r, g in geography.mapping.items() if g == group]
            total = sum(outputs[r] for r in members)
            available = [r for r in members if r in profiles]
            new_outputs.append(dict(Sector=sector, Region=group, Quantity=total, Unit=next(iter(output_units)),
                                    Source=f"Output-weighted aggregation: {', '.join(members)}"))
            if not available:
                continue
            base = settings[available[0]].copy()
            for region in available:
                for column in ("Unit", "Parent Sector", "Add or Split", "Final consumption", "Consumption category", "Leave empty"):
                    if settings[region].get(column) != base.get(column):
                        raise ValueError(f"Inconsistent {column} for {sector}/{group}")
            acc, units, labels = defaultdict(float), {}, {}
            # Zero-output profiles are retained explicitly for future editing,
            # never inferred from a blank production observation.
            if total == 0:
                warnings.warn(f"{sector}/{group}: zero output; using an unweighted reference profile.")
            for region in available:
                weight = outputs[region] / total if total else 1 / len(available)
                for row in profiles[region]:
                    origin = row.get("DB Region")
                    if origin in geography.mapping:
                        origin = geography.mapping[origin]
                    elif origin in mapped_clusters:
                        if mapped_clusters[origin] is None:
                            raise ValueError(f"Origin cluster {origin} only partly covers a target region; resolve origin flows first.")
                    elif origin not in (None, ""):
                        raise ValueError(f"Unmapped input origin: {origin}")
                    if row.get("Change type") not in (None, "Update"):
                        raise ValueError("Only resolved Update profiles can be aggregated.")
                    key = (row.get("Item type"), row.get("DB Item"), origin, row.get("Change type"))
                    if key in units and units[key] != row.get("Unit"):
                        raise ValueError(f"Inconsistent units for {key}")
                    units[key] = row.get("Unit")
                    labels.setdefault(key, row.get("Input"))
                    acc[key] += weight * _number(row.get("Quantity"), f"{sector}/{region}/{key}")
            sheet = f"INV_{len(inventory_sheets) + 1:03d}"
            rows = [dict(Quantity=q, Unit=units[k], Input=labels[k], **{
                "Item type": k[0], "DB Item": k[1], "DB Region": k[2], "Change type": k[3]})
                    for k, q in acc.items()]
            base.update({"Region": group, "Inventory sheet": sheet})
            new_master.append(base)
            inventory_sheets.append((sheet, rows))

        trades = defaultdict(float)
        trade_columns = [v.value for v in book["Trades"][1]] if "Trades" in book else [
            "Sector_from", "Region_from", "Region_to", "Quantity", "Unit", "Source", "Notes"]
        if "Trades" in book:
            for row in _rows(book["Trades"]):
                origin, destination = row["Region_from"], row["Region_to"]
                if origin not in geography.mapping or destination not in geography.mapping:
                    raise ValueError("Trade contains unmapped regions.")
                a, b = geography.mapping[origin], geography.mapping[destination]
                if a != b:
                    trades[(row["Sector_from"], a, b, row["Unit"])] += _number(row["Quantity"], "trade")
        for name in {r["Inventory sheet"] for r in master}:
            del book[name]
        _write(book, "Master", new_master, master_columns)
        _write(book, "Total outputs", new_outputs, ["Sector", "Region", "Quantity", "Unit", "Source", "Notes"])
        _write(book, "Trades", [dict(zip(trade_columns[:5], (*k[:3], q, k[3]))) for k, q in trades.items()], trade_columns)
        # Preserve only expressible cluster definitions. A used partial cluster
        # has already raised above; unused source clusters need not survive.
        _write(book, "Regions Clusters", [
            {name: members[i] if i < len(members) else None for name, members in mapped_clusters.items() if members is not None}
            for i in range(max((len(v) for v in mapped_clusters.values() if v is not None), default=0))],
            [name for name, values in mapped_clusters.items() if values is not None])
        for name, rows in inventory_sheets:
            _write(book, name, rows, ["Quantity", "Unit", "Input", "Item type", "DB Item", "DB Region", "Change type"])
        if "Exclusions" in book and _rows(book["Exclusions"]):
            raise ValueError("Nonempty native Exclusions require an explicit aggregation policy.")
        if "Exclusions" not in book:
            _write(book, "Exclusions", [], ["Region_from", "Region_to", "Sector_from", "Sector_to", "Notes"])
        if "Tolerances" not in book:
            _write(book, "Tolerances", [{"tol_Name": "delta", "values": 0.0}, {"tol_Name": "eps", "values": 0.0}], ["tol_Name", "values"])
        if "Summary" in book:
            ws = book["Summary"]
            for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
                row[0].value, row[0].hyperlink = None, None
            for i, name in enumerate((n for n in book.sheetnames if n != "Summary"), 2):
                ws.cell(i, 9, name).hyperlink = f"#'{name}'!A1"
        target.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=target.parent, delete=False) as stream:
            temporary = Path(stream.name)
        try:
            book.save(temporary)
            temporary.replace(target)
        finally:
            temporary.unlink(missing_ok=True)
        return {"sector": sector, "regions": len(geography.regions), "output": sum(outputs.values()),
                "source": str(source), "source_sha256": hashlib.sha256(raw).hexdigest(),
                "interregional_trade": sum(trades.values()),
                "supplemented_output_regions": supplemented,
                "output_reference": output_reference.get('source') if supplemented else None,
                "output_reference_sha256": output_reference.get('sha256') if supplemented else None}
    finally:
        book.close()
