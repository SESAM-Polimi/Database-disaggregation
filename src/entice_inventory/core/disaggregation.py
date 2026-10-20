"""Prepare a shared geography and one editable MARIO split/step workbook."""
from __future__ import annotations

from dataclasses import replace
import hashlib
import json
from io import BytesIO
from pathlib import Path
import tempfile

from entice_inventory.core.baseline import load_baseline
from entice_inventory.core.geography import Geography
from entice_inventory.inventory.regionalize import aggregate_inventory


def prepare_inputs(source_dir, output_dir, *, geography=None, steps=None,
                   intermediate_steps=False, baseline_path=None, baseline_format=None,
                   output_reference_dir=None):
    """Prepare without solving or altering the canonical native source files.

    ``source_dir`` contains resolved MARIO/D2.4 inventories (the same boundary
    used by Camilla's regional aggregation). The original generation pipeline
    can pass its freshly exported native inventory directory here.
    """
    from mario.ops.add_sector_steps import validate_plan, write_workbook_payload
    if steps is not None and (not isinstance(steps, dict) or any(
            not isinstance(name, str) or not name.strip() or not isinstance(values, list)
            or any(not isinstance(s, str) or not s.strip() for s in values)
            for name, values in steps.items())):
        raise ValueError('steps must map non-empty group names to sector-name lists.')
    source_dir, output_dir = Path(source_dir), Path(output_dir)
    if output_dir.exists():
        raise FileExistsError(f"Use a new preparation directory: {output_dir}")
    files = sorted(p for p in source_dir.glob('*.xlsx') if not p.name.startswith('~$'))
    if not files:
        raise FileNotFoundError(f"No resolved inventory workbooks found in {source_dir}")
    database = load_baseline(baseline_path, baseline_format, geography=geography)
    if geography is None:
        geography = Geography({r: r for r in database.get_index('Region')}, 'native')
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f'.{output_dir.name}_', dir=output_dir.parent))
    inventories = staging / 'inventories'
    results = []
    try:
        reference_dir = Path(output_reference_dir) if output_reference_dir else None
        references = None
        for source in files:
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(source.read_bytes()), read_only=True, data_only=True)
            try:
                sector = next(book['Master'].iter_rows(min_row=2, values_only=True))[1]
                header, *values = book['Total outputs'].values
                native_regions = {dict(zip(header, row))['Region'] for row in values if row[0] is not None}
            finally:
                book.close()
            if geography.mapping.keys() - native_regions and references is None and reference_dir and reference_dir.is_dir():
                references = read_output_references(reference_dir)
            results.append(aggregate_inventory(source, inventories / source.name, geography,
                                               output_reference=(references or {}).get(sector)))
        database.read_add_sectors_excel(inventories, read_inventories=True, split=True)
        workbook = database.add_sectors_workbook
        expected = {r['sector'] for r in results}
        actual = set(workbook.master_sheet['Sector'])
        if expected != actual:
            raise ValueError(f"Native MARIO reader did not retain all inventories: expected={expected}, actual={actual}")
        if intermediate_steps and steps is None:
            steps = {'Step 1': workbook.master_sheet['Sector'].drop_duplicates().tolist()}
        workbook = replace(workbook, intermediate_steps=steps)
        if steps is not None:
            validate_plan(workbook, database.get_index('Sector'))
        write_workbook_payload(workbook, staging / 'add_sectors.xlsx', instance=database)
        database.to_parquet(path=staging / 'baseline', flows=True, coefficients=False, include_meta=True)
        manifest = {'geography': {'name': geography.name, 'mapping': geography.mapping,
                                   'regions': geography.regions}, 'inventories': results,
                    'steps': steps, 'source_files': {r['source']: r['source_sha256'] for r in results},
                    'output_reference_files': {r['output_reference']: r['output_reference_sha256']
                                               for r in results if r['output_reference']},
                    'input_kind': 'resolved_native_profiles', 'status': 'prepared',
                    'zero_output_profile_policy': 'unweighted_reference_with_warning',
                    'support_defaults': 'Missing Exclusions empty; missing delta/eps zero. Existing values retained.'}
        (staging / 'preparation.json').write_text(json.dumps(manifest, indent=2, allow_nan=False))
        staging.rename(output_dir)
    except Exception as exc:
        (staging / 'failure.json').write_text(json.dumps({'status': 'failed', 'error': str(exc)}, indent=2))
        raise
    return output_dir


def read_output_references(directory):
    """Use explicitly stored native outputs to restore sparse export coverage.

    D2.4 export retains only countries with a rendered profile. Original MARIO
    inputs retain their separate output observations, including explicit zeros.
    No absent observation is converted to zero here.
    """
    from openpyxl import load_workbook
    references = {}
    for path in sorted(Path(directory).glob('*.xlsx')):
        if path.name.startswith('~$'):
            continue
        raw = path.read_bytes()
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            if 'Total outputs' not in book:
                continue
            header, *values = book['Total outputs'].values
            rows = [{k: v for k, v in dict(zip(header, row)).items() if k in ('Sector', 'Region', 'Quantity', 'Unit')}
                    for row in values if row[0] is not None]
            sectors = {row['Sector'] for row in rows}
            if len(sectors) != 1:
                raise ValueError(f'Expected one output-reference sector in {path}')
            sector = sectors.pop()
            if sector in references and references[sector]['rows'] != rows:
                raise ValueError(f'Conflicting native output reference files for {sector}')
            references[sector] = {'source': str(path), 'sha256': hashlib.sha256(raw).hexdigest(), 'rows': rows}
        finally:
            book.close()
    return references
