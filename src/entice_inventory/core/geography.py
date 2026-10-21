"""Explicit regional partitions shared by the baseline and resolved inventories.

Aggregation operates on flows. The parquet reader projects column batches, so
preparing a small geography does not require loading the native global Z.
"""
from __future__ import annotations

from dataclasses import dataclass
import csv
import json
from io import BytesIO
from pathlib import Path

from entice_inventory.core.paths import PROJECT_ROOT, local_config, project_data_root


@dataclass(frozen=True)
class Geography:
    mapping: dict[str, str]
    name: str = "custom"

    def __post_init__(self):
        if not self.mapping or any(not isinstance(k, str) or not k.strip() or
                                   not isinstance(v, str) or not v.strip()
                                   for k, v in self.mapping.items()):
            raise ValueError("Geography requires non-empty source and target region codes.")
        if any(k != k.strip() or v != v.strip() for k, v in self.mapping.items()):
            raise ValueError("Region codes must not contain leading/trailing whitespace.")

    @property
    def regions(self):
        return list(dict.fromkeys(self.mapping.values()))

    def validate(self, source_regions):
        source = set(source_regions)
        missing, extra = source - self.mapping.keys(), self.mapping.keys() - source
        if missing or extra:
            raise ValueError(f"Geography does not match the baseline: missing={sorted(missing)}, extra={sorted(extra)}")

    @classmethod
    def read(cls, path, *, name=None, sheet=0):
        path = Path(path)
        if path.suffix.lower() == ".json":
            # JSON pairs retain duplicate keys so an accidental reassignment is rejected.
            rows = json.loads(path.read_text(), object_pairs_hook=lambda pairs: pairs)
        elif path.suffix.lower() in {".csv", ".tsv"}:
            with path.open(newline="", encoding="utf-8-sig") as stream:
                rows = list(csv.reader(stream, delimiter="\t" if path.suffix.lower() == ".tsv" else ","))[1:]
        elif path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=True)
            try:
                ws = book.worksheets[sheet] if isinstance(sheet, int) else book[sheet]
                rows = list(ws.iter_rows(min_row=2, values_only=True))
            finally:
                book.close()
        else:
            raise ValueError("Region mapping must be JSON, CSV, TSV or XLSX.")
        mapping = {}
        for row in rows:
            if not row or all(value is None or value == "" for value in row):
                continue
            if len(row) != 2 or any(not isinstance(v, str) or not v.strip() for v in row):
                raise ValueError(f"Expected two region codes, found {row!r}")
            source, target = (v.strip() for v in row)
            if source in mapping:
                raise ValueError(f"Duplicate source region in mapping: {source}")
            mapping[source] = target
        return cls(mapping, name or path.stem)


def configured_geography():
    config = local_config().get("geography")
    if config is None:
        return None
    if not isinstance(config, dict) or not config.get("mapping"):
        raise ValueError("geography must contain a mapping file path.")
    mapping_base = config.get("mapping_base", "data_root")
    if mapping_base not in {"data_root", "repository"}:
        raise ValueError("geography.mapping_base must be 'data_root' or 'repository'.")
    path = Path(config["mapping"]).expanduser()
    if not path.is_absolute():
        base = PROJECT_ROOT if mapping_base == "repository" else project_data_root(required=True)
        path = base / path
    return Geography.read(path, name=config.get("name"), sheet=config.get("sheet", 0))


def map_axis(axis, geography):
    import pandas as pd
    if not isinstance(axis, pd.MultiIndex) or "Region" not in axis.names:
        return axis
    position = axis.names.index("Region")
    values = axis.get_level_values(position)
    unknown = set(values) - geography.mapping.keys()
    if unknown:
        raise ValueError(f"Unmapped regions on matrix axis: {sorted(unknown)}")
    arrays = [axis.get_level_values(i) if i != position else values.map(geography.mapping)
              for i in range(axis.nlevels)]
    return pd.MultiIndex.from_arrays(arrays, names=axis.names)


def aggregate_frame(frame, geography):
    """Sum only axes that actually carry a Region coordinate."""
    import pandas as pd
    frame = frame.copy()
    for axis in (0, 1):
        labels = frame.index if axis == 0 else frame.columns
        if not isinstance(labels, pd.MultiIndex) or "Region" not in labels.names:
            continue
        labels = map_axis(labels, geography)
        if axis == 0:
            frame.index = labels
            frame = frame.groupby(level=list(range(labels.nlevels)), sort=False).sum(min_count=1)
        else:
            frame.columns = labels
            frame = frame.T.groupby(level=list(range(labels.nlevels)), sort=False).sum(min_count=1).T
    return frame


def aggregate_parquet(path, geography, *, chunk_size=256):
    import numpy as np
    import pyarrow.parquet as pq
    if chunk_size < 1:
        raise ValueError("chunk_size must be positive.")
    parquet = pq.ParquetFile(path)
    metadata = json.loads(parquet.schema_arrow.metadata[b"pandas"])
    index_columns = [v for v in metadata["index_columns"] if isinstance(v, str)]
    columns = [v["field_name"] for v in metadata["columns"] if v["field_name"] not in index_columns]
    result = None
    for offset in range(0, len(columns), chunk_size):
        block = parquet.read(columns=columns[offset:offset + chunk_size] + index_columns).to_pandas()
        if not np.isfinite(block.to_numpy(dtype=float)).all():
            raise ValueError(f"Non-finite baseline flows in {path}")
        block = aggregate_frame(block, geography)
        result = block if result is None else result.add(block, fill_value=0)
    if result is None:
        return parquet.read().to_pandas()
    return result.sort_index().sort_index(axis=1)


def load_regional_baseline(path, geography, *, chunk_size=256):
    import mario
    import pandas as pd
    import pyarrow.parquet as pq
    path = Path(path)
    x_path = path / "X.parquet"
    if x_path.exists():
        source_regions = pd.read_parquet(x_path).index.get_level_values("Region").unique()
    else:
        # Read only the index, without materialising Z's data columns.
        p = pq.ParquetFile(path / "Z.parquet")
        meta = json.loads(p.schema_arrow.metadata[b"pandas"])
        source_regions = p.read(columns=meta["index_columns"]).to_pandas().index.get_level_values("Region").unique()
    geography.validate(source_regions)
    matrices = {name: aggregate_parquet(path / f"{name}.parquet", geography, chunk_size=chunk_size)
                for name in ("Z", "Y", "V", "E", "EY", "VY") if (path / f"{name}.parquet").exists()}
    unit_frame = pd.read_parquet(path / "units.parquet")
    units = {kind: part.droplevel(0) for kind, part in unit_frame.groupby(level=0, sort=False)}
    return mario.Database(name=f"GTAP — {geography.name}", table="IOT", units=units, **matrices)
