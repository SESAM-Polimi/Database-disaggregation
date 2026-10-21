"""Explicit, opt-in reconciliation of factor priors on immutable input copies.

This is a prior transformation, not a correction to measured value added. It
preserves the net factor total and does not alter intermediate inputs or targets.
The ordinary preparation policy remains unchanged until this assumption is reviewed.
"""
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import posixpath
import re
import shutil
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook


def project_factors(prior, nonnegative):
    """Minimize sum((v-q)**2 / abs(q)), preserving sum(v) with sign bounds.

    Zero priors use the smallest nonzero absolute prior as their variance. This
    is recorded by the caller, introduces no fixed monetary floor and leaves an
    already admissible profile exactly unchanged. Tax/TTM can retain either sign.
    The active set solves the diagonal strictly convex problem analytically.
    """
    q = np.asarray(prior, dtype=float)
    bounded = np.asarray(nonnegative, dtype=bool)
    if q.ndim != 1 or not q.size or q.shape != bounded.shape or not np.isfinite(q).all():
        raise ValueError('Expected equally sized, finite factor values and sign mask.')
    if not np.any(bounded & (q < 0)):
        return q.copy()
    target = math.fsum(q)
    if bounded.all() and target < 0:
        raise ValueError('Negative net factor total cannot have only nonnegative factors.')
    nonzero = np.abs(q[q != 0])
    variance = np.where(q != 0, np.abs(q), nonzero.min())
    fixed = np.zeros(q.size, dtype=bool)
    for _ in range(q.size + 1):
        free = ~fixed
        result = np.zeros_like(q)
        if not free.any():
            if target != 0:
                raise ValueError('No free factor can preserve the requested total.')
            return result
        multiplier = (math.fsum(q[free]) - target) / math.fsum(variance[free])
        result[free] = q[free] - multiplier * variance[free]
        violates = bounded & free & (result < 0)
        if violates.any():
            fixed |= violates
            continue
        # Correct roundoff on an unrestricted factor if available; otherwise on
        # the largest positive free factor. This is not a bound relaxation.
        candidates = np.flatnonzero(free & ~bounded)
        anchor = int(candidates[np.argmax(variance[candidates])]) if candidates.size else int(np.argmax(result))
        result[anchor] += target - math.fsum(result)
        if np.any(result[bounded] < 0) or abs(math.fsum(result) - target) > 1e-14 * max(1., abs(target)):
            raise ArithmeticError('Factor projection failed its conservation/domain checks.')
        return result
    raise ArithmeticError('Factor projection did not converge.')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def fingerprints(root):
    return {str(p.relative_to(root)): digest(p) for p in sorted(Path(root).rglob('*')) if p.is_file()}


def _sheet_paths(archive):
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    relationships = {r.attrib['Id']: r.attrib['Target'] for r in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
    rid = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    return {s.attrib['name']: posixpath.normpath(posixpath.join('xl', relationships[s.attrib[rid]]))
            if not relationships[s.attrib[rid]].startswith('/') else relationships[s.attrib[rid]].lstrip('/')
            for s in ET.fromstring(archive.read('xl/workbook.xml')).find('m:sheets', ns)}


def _patch_quantities(payload, replacements):
    """Patch only numeric value tokens; preserve every other byte in the XML."""
    for coordinate, before, after in replacements:
        cell = re.compile(rb'(<c\b[^>]*\br="' + coordinate.encode() + rb'"[^>]*>.*?<v>)([^<]*)(</v>.*?</c>)', re.S)
        matches = list(cell.finditer(payload))
        if len(matches) != 1 or float(matches[0].group(2)) != before:
            raise ValueError(f'Expected exactly one unchanged numeric quantity at {coordinate}.')
        match = matches[0]
        payload = payload[:match.start(2)] + repr(float(after)).encode() + payload[match.end(2):]
    ET.fromstring(payload)
    return payload


def reconcile_prepared_factors(source, output, sectors, signed_factors=('Tax', 'TTM')):
    """Copy preparation; change only selected profiles with invalid factor signs.

    The XLSX is a machine input container with thousands of independent inventory
    sheets. A targeted ZIP/XML update avoids a workbook-wide authoring roundtrip:
    all controls, targets, styles, formulas and unselected sheets stay byte-identical.
    """
    source, output = Path(source).resolve(), Path(output).resolve()
    if not sectors or output.exists() or source == output or source in output.parents:
        raise ValueError('Select sectors and a fresh output outside the immutable preparation.')
    workbook = source/'add_sectors.xlsx'
    source_hash, baseline_hashes = digest(workbook), fingerprints(source/'baseline')
    changes, patches, found = [], {}, set()
    book = load_workbook(workbook, read_only=True, data_only=False)
    try:
        master = book['Master']
        header = next(master.values)
        for values in master.iter_rows(min_row=2, values_only=True):
            row = dict(zip(header, values))
            if row.get('Sector') not in sectors:
                continue
            found.add(row['Sector'])
            sheet = row['Inventory sheet']
            if sheet in patches:
                raise ValueError('A selected inventory sheet is referenced by multiple profiles.')
            table = book[sheet]
            columns = {value: i for i, value in enumerate(next(table.values))}
            factors = []
            for cells in table.iter_rows(min_row=2):
                if cells[columns['Item type']].value == 'Factor of production':
                    cell = cells[columns['Quantity']]
                    factors.append((cells[columns['DB Item']].value, cell.coordinate, float(cell.value)))
            if not factors:
                raise ValueError(f'No factor priors in {sheet}.')
            q = [f[2] for f in factors]
            bounded = [f[0] not in signed_factors for f in factors]
            if not any(b and v < 0 for v, b in zip(q, bounded)):
                continue
            result = project_factors(q, bounded)
            patches[sheet] = [(cell, before, float(after)) for (_, cell, before), after in zip(factors, result) if before != after]
            changes.append(dict(sector=row['Sector'], region=row['Region'], sheet=sheet,
                net_before=math.fsum(q), net_after=math.fsum(result),
                zero_prior_variance=min(abs(v) for v in q if v != 0),
                factors=[dict(factor=factor, cell=cell, before=before, after=float(after),
                              delta=float(after)-before, nonnegative=bounded[i], active_bound=bool(bounded[i] and after == 0))
                         for i, ((factor, cell, before), after) in enumerate(zip(factors, result))]))
    finally:
        book.close()
    if found != set(sectors):
        raise ValueError(f'Unresolved selected sectors: {set(sectors)-found}')
    if not changes:
        raise ValueError('No invalid factor priors found; no transformed preparation was created.')
    output.mkdir(parents=True)
    shutil.copytree(source/'baseline', output/'baseline')
    with ZipFile(workbook) as original, ZipFile(output/'add_sectors.xlsx', 'w') as transformed:
        paths = _sheet_paths(original)
        changed_members = {paths[sheet]: values for sheet, values in patches.items()}
        for member in original.infolist():
            payload = original.read(member.filename)
            if member.filename in changed_members:
                payload = _patch_quantities(payload, changed_members[member.filename])
            transformed.writestr(member, payload)
    with ZipFile(workbook) as original, ZipFile(output/'add_sectors.xlsx') as transformed:
        actual = [p for p in original.namelist() if original.read(p) != transformed.read(p)]
        if set(actual) != set(changed_members):
            raise ArithmeticError('Unexpected XLSX member changes.')
        # Reversing exactly the audited numeric edits must recover the source XML.
        for path, values in changed_members.items():
            if _patch_quantities(transformed.read(path), [(c, a, b) for c, b, a in values]) != original.read(path):
                # Numeric tokens may have different equivalent decimal spellings;
                # compare parsed cells, in addition to the byte checks above.
                before_tree = ET.fromstring(original.read(path)); after_tree = ET.fromstring(transformed.read(path))
                permitted = {c for c, _, _ in values}
                for a, b in zip(before_tree.iter(), after_tree.iter()):
                    if a.tag != b.tag or a.attrib != b.attrib:
                        raise ArithmeticError('Unexpected inventory structure change.')
                original_cells = {c.attrib['r']: ET.tostring(c) for c in before_tree.iter() if c.tag.endswith('}c')}
                new_cells = {c.attrib['r']: ET.tostring(c) for c in after_tree.iter() if c.tag.endswith('}c')}
                if any(original_cells[c] != new_cells[c] for c in original_cells if c not in permitted):
                    raise ArithmeticError('Unexpected inventory cell change.')
    unchanged = (digest(workbook) == source_hash and fingerprints(source/'baseline') == baseline_hashes
                 and fingerprints(output/'baseline') == baseline_hashes)
    if not unchanged:
        raise ArithmeticError('A source input or copied baseline changed.')
    audit = dict(format='factor-domain-reconciliation-v1', created_at=datetime.now(timezone.utc).isoformat(),
        source_preparation=str(source), source_workbook_sha256=source_hash,
        reconciled_workbook_sha256=digest(output/'add_sectors.xlsx'), source_baseline_sha256=baseline_hashes,
        selected_sectors=sorted(sectors), signed_factors=list(signed_factors),
        rule='Minimize sum((v-q)^2/abs(q)); zero-prior variance uses smallest nonzero abs(q); preserve sum(v); bound all factors except Tax/TTM at zero.',
        interpretation='Opt-in prior assumption, not new empirical evidence. No intermediate-input, output, trade, tolerance, exclusion, plan or baseline changes.',
        profiles=changes, changed_zip_members=actual, unchanged_zip_members_verified=True, sources_unchanged=True)
    (output/'reconciliation.json').write_text(json.dumps(audit, indent=2, allow_nan=False))
    if (source/'preparation.json').is_file():
        preparation = json.loads((source/'preparation.json').read_text())
        preparation['factor_domain_reconciliation'] = dict(source=str(source), audit='reconciliation.json',
            original_preparation_sha256=digest(source/'preparation.json'))
        (output/'preparation.json').write_text(json.dumps(preparation, indent=2, allow_nan=False))
    return audit
