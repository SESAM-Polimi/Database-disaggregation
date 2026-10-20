"""Explicit source columns, sparse-domain zeros and unambiguous geography."""
from collections import defaultdict
from functools import lru_cache
import math

from openpyxl.utils.cell import column_index_from_string


def number(value, *, context):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        raise ValueError(f'Boolean monetary value: {context}')
    result = float(value)
    if not math.isfinite(result) or result < 0:
        raise ValueError(f'Invalid nonnegative source value: {context}')
    return result


class Sources:
    def __init__(self, paths, registry, native, policy, audit):
        self.paths, self.registry, self.native, self.policy, self.audit = paths, registry, native, policy, audit
        self.region_lookup = {str(v).strip().casefold():r['id'] for r in registry.regions()
                              for v in (r['id'], r['name']) if r['id'] in native.regions}
        self.paired_trade, self.customs_trade = {}, {}
        self.domain_sectors, self.domain_regions = set(), set()
        self.purdue_output = {}
        if any(r['output_source']['kind']=='purdue' for r in policy['sectors'].values()):
            self._read_purdue()
        if paths.get('purdue_trade') is not None:
            self._read_customs()

    def _read_purdue(self):
        path = self.paths['purdue_split_targets']
        book = self.audit.workbook(path, role='paired_output_trade')
        try:
            for name, target in (('REG',self.domain_regions),('REPSEC',self.domain_sectors)):
                for row in list(book[name].values)[1:]:
                    if row[0] is not None:
                        target.add(str(row[0]).strip().upper())
            if set(self.native.regions) - self.domain_regions:
                raise ValueError('Purdue REG does not cover the native baseline geography.')
            seen = set()
            for i,row in enumerate(book['REPOUT'].iter_rows(min_row=2,values_only=True),2):
                region, sector = str(row[0]).strip().upper(), str(row[1]).strip().upper()
                if region not in self.domain_regions or sector not in self.domain_sectors:
                    raise ValueError(f'Out-of-domain Purdue REPOUT row {i}')
                key = sector, region
                if key in seen:
                    raise ValueError(f'Duplicate Purdue REPOUT key: {key}')
                seen.add(key)
                value = number(row[2],context=f'{path}:REPOUT!C{i}')
                self.purdue_output[key] = (0. if value is None else value, f'{path}:REPOUT!C{i}',
                                           'stored_zero' if value == 0 else 'blank_parameter_zero' if value is None else 'numeric')
            trade = defaultdict(dict)
            for i,row in enumerate(book['BLTTRD'].iter_rows(min_row=2,values_only=True),2):
                origin,destination,sector = (str(v).strip().upper() for v in row[:3])
                value = number(row[3],context=f'{path}:BLTTRD!D{i}')
                if origin not in self.domain_regions or destination not in self.domain_regions or sector not in self.domain_sectors:
                    raise ValueError(f'Out-of-domain Purdue BLTTRD row {i}')
                if (origin,destination) in trade[sector]:
                    raise ValueError(f'Duplicate Purdue trade coordinate: {sector}/{origin}/{destination}')
                trade[sector][(origin,destination)] = value or 0.
            self.paired_trade = dict(trade)
        finally:
            book.close()

    def _read_customs(self):
        path = self.paths['purdue_trade']
        book = self.audit.workbook(path,role='customs_trade')
        try:
            tables = defaultdict(dict)
            for i,row in enumerate(book['NTSCIF'].iter_rows(min_row=2,values_only=True),2):
                sector,origin,destination = (str(v).strip().upper() for v in row[:3])
                if origin not in self.native.regions or destination not in self.native.regions:
                    continue
                value = number(row[3],context=f'{path}:NTSCIF!D{i}')
                if (origin,destination) in tables[sector]:
                    raise ValueError(f'Duplicate NTSCIF coordinate: {sector}/{origin}/{destination}')
                tables[sector][(origin,destination)] = value or 0.
            self.customs_trade = dict(tables)
        finally:
            book.close()

    def _clusters(self, book):
        groups = defaultdict(set)
        current = None
        for i,row in enumerate(book['Region'].iter_rows(values_only=True),1):
            identifier, member = (row[0] if len(row)>0 else None), (row[1] if len(row)>1 else None)
            code = self.region_lookup.get(str(member).strip().casefold())
            if 3 <= i <= 165 and identifier and code:
                groups[str(identifier).strip()].add(code)
            elif i >= 328:
                if identifier:
                    current = str(identifier).strip()
                if current and code:
                    groups[current].add(code)
        groups['GLOBAL'] = set(self.native.regions)
        groups['WORLD'] = set(self.native.regions)
        return groups

    def observations(self, sector, parent):
        rule = self.policy['sectors'][sector]['output_source']
        kind = rule['kind']
        if kind == 'legacy_share':
            return []
        if kind == 'purdue':
            if sector not in self.domain_sectors:
                raise ValueError(f'{sector} is absent from Purdue REPSEC; no implicit zeros are valid.')
            result = []
            for region in self.native.regions:
                value, source, state = self.purdue_output.get((sector,region),
                    (0.,f"{self.paths['purdue_split_targets']}:REPOUT[{sector},{region}]",'implicit_parameter_zero'))
                result.append({'members':[region],'value':value,'identifier':region,'source':source,
                               'state':state,'kind':'direct','confidence':1.})
                self.audit.record('paired_parameter',kind='source',sector=sector,region=region,
                    before=value,after=value,source=source,state=state,unit='M USD',assumption=rule['assumption'])
            return result
        path = self.paths['partner_templates']/rule['file']
        book = self.audit.workbook(path,role='partner_template')
        formulas = self.audit.workbook(path,role='partner_template',data_only=False)
        try:
            groups = self._clusters(book)
            values = list(book['Total production'].values)
            formula_rows = list(formulas['Total production'].values)
            col = column_index_from_string(rule['column'])-1
            headers=[r[col] if col<len(r) else None for r in values[:2]]
            if rule.get('expected_headers') is not None and headers!=rule['expected_headers']:
                raise ValueError(f'{sector}: selected source column headers changed; revise the policy: {headers}')
            if 'General info' in book:
                info=list(book['General info'].values)
                self.audit.record('template_identity_context',kind='identity',sector=sector,source=str(path),
                    declared_activity=info[4][3],declared_parent=info[6][3],
                    assumption='Explicit policy file-to-sector association; copied General info is diagnostic, not an automatic alias')
            unit_col = column_index_from_string(rule['unit_column'])-1 if rule.get('unit_column') else None
            candidates = defaultdict(list)
            for i,row in enumerate(values[2:],3):
                if not row or row[0] is None:
                    continue
                identifier = str(row[0]).strip()
                raw = row[col] if col < len(row) else None
                region = self.region_lookup.get(identifier.casefold())
                proxy=self.policy['geography'].get('region_proxies',{}).get(sector,{}).get(identifier)
                if proxy:
                    if kind!='template_share' or proxy['target'] not in self.native.regions:
                        raise ValueError('Geography share proxies require a template share and a native residual region.')
                    region=proxy['target']
                members = [region] if region else sorted(groups.get(identifier,[]))
                if not members:
                    try:
                        members = self.registry.members(identifier)
                    except LookupError:
                        if raw not in (None,'',0,0.):
                            raise ValueError(f'{sector}: positive output has unknown geography {identifier!r}')
                        continue
                try:
                    raw_number = number(raw,context=f'{path}:Total production!{rule["column"]}{i}')
                    invalid_value=False
                except (ValueError,TypeError):
                    raw_number=None;invalid_value=True
                formula = formula_rows[i-1][col] if col < len(formula_rows[i-1]) else None
                formula = formula if isinstance(formula,str) and formula.startswith('=') else None
                source = f'{path}:Total production!{rule["column"]}{i}'
                converted = None if raw_number is None else raw_number*float(rule.get('factor',1.))
                if kind == 'template_share' and converted is not None:
                    if len(members) != 1:
                        raise ValueError('Parent-share observations must identify one producing region.')
                    converted *= float(self.native.budgets.output.loc[(members[0],parent)])
                state = 'unusable_value' if invalid_value else 'uncached_formula' if raw_number is None and formula else 'blank' if raw_number is None else 'stored_zero' if raw_number==0 else 'numeric'
                self.audit.record('explicit_column',kind='source',sector=sector,identifier=identifier,
                    before=raw_number,after=converted,source=source,state=state,formula=formula,
                    raw_text=str(raw) if invalid_value else None,
                    column_headers=headers,target_year=self.policy['target_year'],geography_proxy=proxy,
                    raw_unit=row[unit_col] if unit_col is not None and unit_col<len(row) else None,
                    factor=float(rule.get('factor',1)),assumption=rule['assumption'])
                if converted is not None:
                    candidates[identifier].append({'members':members,'value':converted,'identifier':identifier,
                        'source':source,'state':state,'kind':'direct' if region else 'cluster',
                        'confidence':proxy['confidence'] if proxy else 0.5 if rule.get('factor',1)!=1 or kind=='template_share' else 1.})
            result=[]
            for identifier,rows in candidates.items():
                positive = [r for r in rows if r['value']>0]
                selected = positive or rows
                if len({r['value'] for r in selected})>1:
                    raise ValueError(f'{sector}: conflicting repeated output observations for {identifier}')
                result.append(selected[0])
                if len(rows)>1:
                    self.audit.record('duplicate_observation',kind='source',sector=sector,identifier=identifier,
                        selected=selected[0]['source'],sources=[r['source'] for r in rows],
                        assumption='One identical positive observation preferred over repeated blanks/zeros; no summation')
            return result
        finally:
            book.close()
            formulas.close()
