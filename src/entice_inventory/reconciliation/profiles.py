"""Donor technologies, origin projection and signed monetary balancing."""
from collections import defaultdict
import hashlib
import math

import numpy as np
from openpyxl import Workbook


INVENTORY_COLUMNS=['Quantity','Unit','Input','Item type','DB Item','DB Region','Change type']
MASTER_COLUMNS=['Region','Sector','Inventory sheet','Quantity','Unit','Final consumption',
                'Consumption category','Parent Sector','Leave empty','Source','Notes','Add or Split']


def rows(book,name):
    header,*values=book[name].values
    return [dict(zip(header,row)) for row in values if any(v is not None for v in row)]


def clusters(book,name):
    if name not in book:return {}
    return {column[0]:[v for v in column[1:] if v is not None]
            for column in zip(*book[name].values) if column[0] is not None}


def expand(item,mapping,universe,trail=()):
    if item in universe:return [item]
    if item in trail or item not in mapping:
        raise ValueError(f'Unknown or cyclic inventory item {item!r}')
    return list(dict.fromkeys(v for child in mapping[item]
                             for v in expand(child,mapping,universe,(*trail,item))))


def balance_profile(profile,policy,audit,*,sector,region):
    """Project positive monetary priors; negative factors and satellites stay fixed."""
    profile=[dict(row) for row in profile]
    monetary=[i for i,row in enumerate(profile) if row['Item type'] in ('Sector','Factor of production')]
    if not monetary:raise ValueError(f'{sector}/{region}: no monetary technology prior.')
    for i in monetary:
        if profile[i]['Item type']=='Sector' and profile[i]['Quantity']<0:
            audit.record('nonnegative_input_projection',kind='profile',sector=sector,region=region,
                         item=profile[i]['DB Item'],before=profile[i]['Quantity'],after=0.)
            profile[i]['Quantity']=0.
    fixed=[i for i in monetary if profile[i]['Quantity']<0]
    positive=[i for i in monetary if profile[i]['Quantity']>0]
    if not positive:raise ValueError(f'{sector}/{region}: no positive monetary component.')
    target=1.-sum(profile[i]['Quantity'] for i in fixed)
    prior=np.array([profile[i]['Quantity'] for i in positive],dtype=float)
    old_total=float(sum(profile[i]['Quantity'] for i in monetary))
    if abs(old_total-1.)<=policy['profiles']['small_residual']:
        result=prior*target/float(prior.sum())
        rule='small_monetary_normalization'
    else:
        weights=np.array([policy['profiles']['sector_weight'] if profile[i]['Item type']=='Sector'
                          else policy['profiles']['factor_weight'] for i in positive])
        metric=prior/weights
        active=np.ones(len(prior),dtype=bool);result=np.zeros(len(prior))
        for _ in range(len(prior)+1):
            step=(target-float(prior[active].sum()))/float(metric[active].sum())
            result[active]=prior[active]+step*metric[active]
            negative=active & (result<0)
            if not negative.any():break
            active[negative]=False;result[negative]=0.
        rule='weighted_monetary_projection'
    for i,value in zip(positive,result):
        before=profile[i]['Quantity'];profile[i]['Quantity']=float(value)
        if before!=float(value):
            audit.record(rule,kind='profile',sector=sector,region=region,item=profile[i]['DB Item'],
                item_type=profile[i]['Item type'],origin=profile[i]['DB Region'],before=before,after=float(value),
                assumption='Weighted chi-square distance; existing negative factors fixed, satellites outside monetary identity')
    total=sum(row['Quantity'] for row in profile if row['Item type'] in ('Sector','Factor of production'))
    if abs(total-1.)>1e-10 or not all(math.isfinite(row['Quantity']) for row in profile):
        raise ValueError(f'{sector}/{region}: monetary balancing failed.')
    return profile


class ProfileBuilder:
    def __init__(self,native,geography,parents,native_output,policy,audit):
        self.native,self.geo,self.parents,self.output,self.policy,self.audit=native,geography,parents,native_output,policy,audit
        self.origin_clusters={'GLOBAL':geography.regions}

    def canonical(self,profile,producing_region,parent,item_clusters,factor_clusters,region_clusters):
        result=[]
        for raw in profile:
            row=dict(raw);kind=row.get('Item type')
            if kind not in ('Sector','Factor of production','Satellite account'):
                raise ValueError(f'Unresolved inventory item type: {kind}')
            if row.get('Change type') not in (None,'Update'):
                raise ValueError('Cleaning expects resolved Update profiles, not partial overrides.')
            value=float(row['Quantity'])
            if not math.isfinite(value):raise ValueError('Non-finite technology observation.')
            if kind=='Sector':
                items=expand(row['DB Item'],item_clusters,set(self.native.sectors)|self.parents.keys())
            elif kind=='Factor of production':
                items=expand(row['DB Item'],factor_clusters,set(self.native.units[kind].index))
            else:
                items=expand(row['DB Item'],{},set(self.native.units[kind].index))
            weights={item:1. for item in items}
            if len(items)>1:
                base,_=self.native.profile((producing_region,),parent)
                base_values={(r['Item type'],r['DB Item']):r['Quantity'] for r in base}
                weights={item:base_values.get((kind,self.parents.get(item,item)),0.) for item in items}
                if kind=='Sector':weights={item:max(0.,q) for item,q in weights.items()}
                if abs(sum(weights.values()))<1e-15:weights={item:1. for item in items}
                total=sum(weights.values());weights={item:q/total for item,q in weights.items()}
            for item,fraction in weights.items():
                target_unit=self.native.units[kind].loc[self.parents.get(item,item) if kind=='Sector' else item,'unit']
                if row.get('Unit')!=target_unit:
                    raise ValueError(f'Profile unit {row.get("Unit")} incompatible with {kind}/{item}: {target_unit}')
                origins={None:1.}
                if kind=='Sector':
                    origin=row.get('DB Region') or 'GLOBAL'
                    if origin in self.geo.mapping:origins={self.geo.mapping[origin]:1.}
                    elif origin=='GLOBAL':origins={'GLOBAL':1.}
                    else:
                        if origin not in region_clusters or not region_clusters[origin]:
                            raise ValueError(f'Unknown inventory origin cluster {origin!r}')
                        members=set(region_clusters[origin])
                        if members-self.geo.mapping.keys():raise ValueError('Unknown member in an inventory origin cluster.')
                        groups=sorted({self.geo.mapping[r] for r in members})
                        full={r for r,g in self.geo.mapping.items() if g in groups}
                        if full==members:
                            name='Origins_'+hashlib.sha256('|'.join(groups).encode()).hexdigest()[:10]
                            self.origin_clusters[name]=groups;origins={name:1.}
                        else:
                            allocation=self.native.origin_weights(producing_region,parent,
                                [self.parents.get(item,item)],sorted(members))
                            origins=defaultdict(float)
                            for r,q in allocation.items():origins[self.geo.mapping[r]]+=q
                            self.audit.record('partial_origin_cluster',kind='profile',producing_region=producing_region,
                                input_item=item,source_cluster=origin,members=sorted(members),weights=dict(origins),
                                assumption='Allocate within partial cluster using parent input-origin pattern; fallback output weights then uniform')
                for origin,origin_share in origins.items():
                    result.append({'Quantity':value*fraction*origin_share,'Unit':target_unit,'Input':item,
                                   'Item type':kind,'DB Item':item,'DB Region':origin,'Change type':'Update'})
        return result

    def build(self,path,sector,old_parent):
        book=self.audit.workbook(path,role='resolved_profiles',read_only=False)
        try:
            master=rows(book,'Master');parent=self.parents[sector]
            region_clusters=clusters(book,'Regions Clusters')
            item_clusters=clusters(book,'Sectors Clusters');factor_clusters=clusters(book,'Factors Clusters')
            raw_profiles={};invalid={}
            if 'Exclusions' in book and rows(book,'Exclusions'):
                raise ValueError('Nonempty native structural exclusions need explicit coordinate decisions before cleaning.')
            for row in master:
                producing=row['Region'];members=[producing] if producing in self.native.regions else region_clusters.get(producing,[])
                if not members:raise ValueError(f'Unresolved producing geography {producing}')
                frame=rows(book,row['Inventory sheet'])
                for region in members:
                    self.audit.record('profile_source_sheet',kind='profile_lineage',sector=sector,region=region,
                                      source=str(path),sheet=row['Inventory sheet'])
                    if region in raw_profiles:raise ValueError(f'Overlapping resolved profiles for {sector}/{region}')
                    try:
                        if not frame or not any(r.get('Item type') in ('Sector','Factor of production') and float(r['Quantity'])>0 for r in frame):
                            raise ValueError('Empty or zero monetary profile')
                        raw_profiles[region]=self.canonical(frame,region,parent,item_clusters,factor_clusters,region_clusters)
                    except (ValueError,TypeError,KeyError) as exc:
                        invalid[region]=str(exc)
            result={}
            for group in self.geo.regions:
                members=[r for r in self.native.regions if self.geo.mapping[r]==group]
                if parent!=old_parent:
                    profile,donors=self.native.profile(tuple(members),parent)
                    self.audit.record('reparented_technology',kind='profile',sector=sector,region=group,
                        old_parent=old_parent,new_parent=parent,donors=list(donors),
                        assumption='Rebuild parent-dependent prior from the selected parent, preserving original child profile as source evidence')
                    result[group]=balance_profile(profile,self.policy,self.audit,sector=sector,region=group)
                    continue
                total=sum(self.output[(sector,r)] for r in members)
                contribution=defaultdict(float)
                for region in members:
                    weight=self.output[(sector,region)]/total if total>0 else 1./len(members)
                    if weight==0:continue
                    if region in raw_profiles:
                        donors=[region];mode='own'
                    else:
                        donors=[r for r in members if r in raw_profiles and self.output[(sector,r)]>0]
                        mode='same_child_macroregion'
                        if not donors:
                            donors=[r for r in self.native.regions if r in raw_profiles and self.output[(sector,r)]>0]
                            mode='same_child_global'
                        if not donors:
                            donors=list(raw_profiles);mode='same_child_unweighted'
                    denominator=sum(self.output[(sector,r)] for r in donors)
                    if donors:
                        weights={r:self.output[(sector,r)]/denominator if denominator>0 else 1./len(donors) for r in donors}
                        for r,q in weights.items():contribution[r]+=weight*q
                    else:
                        contribution[None]+=weight;weights={}
                    self.audit.record('profile_donor',kind='profile_lineage',sector=sector,region=region,
                        selected_region=group,output_weight=weight,donors=weights,mode=mode if donors else 'baseline_parent',
                        invalid_source_reason=invalid.get(region),
                        assumption='Resolved child technologies transferable within macroregion, then globally, then selected parent')
                acc=defaultdict(float);units={}
                for donor,weight in contribution.items():
                    profile=raw_profiles[donor] if donor is not None else self.native.profile(tuple(members),parent)[0]
                    for row in profile:
                        key=(row['Item type'],row['DB Item'],row['DB Region'])
                        if key in units and units[key]!=row['Unit']:raise ValueError('Inconsistent donor units.')
                        units[key]=row['Unit'];acc[key]+=weight*row['Quantity']
                profile=[{'Quantity':q,'Unit':units[k],'Input':k[1],'Item type':k[0],
                          'DB Item':k[1],'DB Region':k[2],'Change type':'Update'} for k,q in acc.items() if q!=0]
                monetary_sum=sum(r['Quantity'] for r in profile if r['Item type']=='Sector')
                if not any(r['Item type']=='Factor of production' for r in profile) and abs(monetary_sum-1.)>self.policy['profiles']['small_residual']:
                    factors=[dict(r) for r in self.native.profile(tuple(members),parent)[0] if r['Item type']=='Factor of production']
                    factor_sum=sum(r['Quantity'] for r in factors)
                    factor_scale=(1.-monetary_sum)/factor_sum if monetary_sum<1 and factor_sum>0 else 1.
                    for row in factors:row['Quantity']*=factor_scale
                    profile.extend(factors)
                    self.audit.record('missing_value_added_prior',kind='profile',sector=sector,region=group,
                        monetary_input_sum=monetary_sum,factor_scale=factor_scale,
                        assumption='Missing entire factor block uses signed parent composition; a positive gap is allocated to that composition')
                result[group]=balance_profile(profile,self.policy,self.audit,sector=sector,region=group)
            return result,master[0]
        finally:
            book.close()

    def write(self,path,sector,profiles,settings,outputs,trade,selected_database):
        from mario.ops.add_sector_workbook import build_db_units_sheet
        book=Workbook();del book['Sheet']
        summary=book.create_sheet('Summary')
        for row in [('INFO',),('Sector',sector,self.policy['sectors'][sector]['name']),
                    ('Parent sector',self.parents[sector]),('Inventory version',self.policy['id']),
                    ('Data collection lead','ENTICE reconciliation; see ../audit/changes.jsonl')]:summary.append(row)
        ws=book.create_sheet('Master');ws.append(MASTER_COLUMNS)
        for i,region in enumerate(self.geo.regions,1):
            row={**settings,'Region':region,'Sector':sector,'Inventory sheet':f'INV_{i:03d}',
                 'Parent Sector':self.parents[sector],'Unit':self.policy['currency_unit'],'Add or Split':'Split',
                 'Source':self.policy['id']}
            ws.append([row.get(c) for c in MASTER_COLUMNS])
            inv=book.create_sheet(row['Inventory sheet']);inv.append(INVENTORY_COLUMNS)
            for value in profiles[region]:inv.append([value.get(c) for c in INVENTORY_COLUMNS])
        for name,mapping in [('Regions Clusters',self.origin_clusters),('Sectors Clusters',{'Cluster1':[]}),
                             ('Factors Clusters',{'Factor1':[]})]:
            ws=book.create_sheet(name);ws.append(list(mapping))
            for i in range(max((len(v) for v in mapping.values()),default=0)):
                ws.append([v[i] if i<len(v) else None for v in mapping.values()])
        ws=book.create_sheet('Total outputs');ws.append(['Sector','Region','Quantity','Unit','Source','Notes'])
        for region in self.geo.regions:ws.append([sector,region,outputs[(sector,region)],self.policy['currency_unit'],self.policy['id'],None])
        ws=book.create_sheet('Trades');ws.append(['Sector_from','Region_from','Region_to','Quantity','Unit','Source','Notes'])
        for (c,o,d),q in sorted(trade.items()):
            if c==sector and q>0:ws.append([c,o,d,q,self.policy['currency_unit'],self.policy['id'],None])
        ws=book.create_sheet('Exclusions');ws.append(['Region_from','Region_to','Sector_from','Sector_to','Notes'])
        ws=book.create_sheet('Tolerances');ws.append(['tol_Name','values']);ws.append(['delta',0.]);ws.append(['eps',0.])
        units=build_db_units_sheet(selected_database)
        ws=book.create_sheet('DB units');ws.append(list(units.columns))
        for row in units.itertuples(index=False,name=None):ws.append(row)
        path.parent.mkdir(parents=True,exist_ok=True);book.save(path);book.close()
