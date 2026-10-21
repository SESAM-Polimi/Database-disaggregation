"""Trace MCU factor priors and export donor proposals without changing inputs.

No model or solver is invoked. Source workbook values are read, never authored.
"""
import argparse
import ast
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import subprocess

from openpyxl import load_workbook
import pandas as pd
import pyarrow.parquet as pq
from entice_inventory.core.paths import data_path
from entice_inventory.reconciliation.prior_review import parent_factor_proposal

VA_CODES={'fixcap','othrent','lablskl','labmskl','labhskl','othtax','prodtax'}


def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''):h.update(chunk)
    return h.hexdigest()


def records(sheet):
    header=next(sheet.values)
    return [dict(zip(header,row)) for row in sheet.iter_rows(min_row=2,values_only=True) if any(v is not None for v in row)]


def summary(rows):
    z=[r for r in rows if r.get('Item type')=='Sector']
    v={r['DB Item']:float(r['Quantity']) for r in rows if r.get('Item type')=='Factor of production'}
    zsum=math.fsum(float(r['Quantity']) for r in z)
    return dict(z_total=zsum,v_total=math.fsum(v.values()),factors=v,
        monetary_total=zsum+math.fsum(v.values()),
        oil_gas_share_of_z=math.fsum(float(r['Quantity']) for r in z if r['DB Item'] in ['OIL','GAS'])/zsum if zsum else None,
        invalid_nonnegative_factors=[key for key,value in v.items() if key not in ['Tax','TTM','OTHTAX'] and value<0])


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--prepared',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    a=p.parse_args();prepared=a.prepared.resolve();out=a.output.resolve()
    if out.exists():p.error('Use a fresh output directory.')
    resolved=Path(json.loads((prepared/'reconciled/complete.json').read_text())['inventory_sources']['MCU'])
    source=data_path('mario_inventories')/'Add_sector_Mining of copper ores.xlsx'
    csv_path=data_path('exioiot');current=prepared/'reconciled/inventories/MCU.xlsx'
    baseline=prepared/'baseline/flows'
    files=[csv_path,source,resolved,current,prepared/'add_sectors.xlsx',prepared/'preparation.json',
           *[baseline/(n+'.parquet') for n in ['X','V','Z','units']]]
    hashes={str(f):sha(f) for f in files}
    raw={}
    with csv_path.open(newline='') as f:
        reader=csv.reader(f);next(reader)
        for line,row in enumerate(reader,2):
            if len(row)>=4 and row[1].lower()=='mcu':
                region=row[2].upper();item=row[0].lower();value=float(row[3])
                if item in raw.setdefault(region,{}):raise ValueError('Duplicate MCU raw coordinate.')
                raw[region][item]=dict(value=value,line=line)
    native=[]
    for region,items in raw.items():
        total=math.fsum(v['value'] for v in items.values());net=math.fsum(v['value'] for k,v in items.items() if k in VA_CODES)
        capital=math.fsum(items.get(k,{'value':0})['value'] for k in ['fixcap','othrent'])
        labor=math.fsum(items.get(k,{'value':0})['value'] for k in ['lablskl','labmskl','labhskl'])
        native.append(dict(region=region,raw_total=total,raw_net_v=net,normalized_net_v=net/total,
            raw_capital=capital,raw_labor=labor,invalid_positive_factor=capital<0 or labor<0))
    book=load_workbook(source,read_only=True,data_only=True)
    try:
        original=records(book['MCU_BRA']);original_summary=summary(original)
        residual_rows=[r for r in original if r.get('Item type')=='Sector' and str(r.get('DB Item','')).startswith('ALL_OTHER')]
        cluster_members=[c[0] for c in list(book['Factors Clusters'].values)[1:] if c[0]]
    finally:book.close()
    book=load_workbook(resolved,read_only=True,data_only=True)
    try:resolved_summary=summary(records(book['MCU_BRA']))
    finally:book.close()
    book=load_workbook(current,read_only=True,data_only=True)
    try:
        master=records(book['Master']);profiles={r['Region']:records(book[r['Inventory sheet']]) for r in master}
        outputs={r['Region']:float(r['Quantity']) for r in records(book['Total outputs'])}
        parent_by_region={r['Region']:r['Parent Sector'] for r in master}
    finally:book.close()
    regional=[]
    for region,profile in profiles.items():
        s=summary(profile)
        regional.append(dict(region=region,output_MUSD=outputs[region],active_output=outputs[region]>1e-6,
            z_total=s['z_total'],v_total=s['v_total'],**s['factors'],
            oil_gas_share_of_z=s['oil_gas_share_of_z'],
            invalid_nonnegative_factors=';'.join(s['invalid_nonnegative_factors'])))
    active_bad=[r for r in regional if r['active_output'] and r['invalid_nonnegative_factors']]
    x=pd.read_parquet(baseline/'X.parquet');v=pd.read_parquet(baseline/'V.parquet')
    parquet=pq.ParquetFile(baseline/'Z.parquet');meta=json.loads(parquet.schema_arrow.metadata[b'pandas'])
    indexes=[n for n in meta['index_columns'] if isinstance(n,str)]
    columns=[c['field_name'] for c in meta['columns'] if c['field_name'] not in indexes]
    proposals=[];coefficient_rows=[];parent_references={}
    for r in active_bad:
        region=r['region'];parent=parent_by_region[region];column=(region,'Sector',parent)
        parent_x=float(x.loc[column].iloc[0])
        if parent_x<=0:raise ValueError('The selected regional parent has zero output.')
        factors=(v.loc[:,column]/parent_x).to_dict()
        selected=[c for c in columns if ast.literal_eval(c)==column]
        if len(selected)!=1:raise ValueError('Ambiguous parent column.')
        z=parquet.read(columns=selected+indexes).to_pandas().iloc[:,0]/parent_x
        z=z.groupby(level='Item').sum()
        parent_references[region]=dict(parent=parent,output_MUSD=parent_x,factors=factors,
            v_total=math.fsum(factors.values()),z_total=float(z.sum()),
            oil_gas_share_of_z=float(z.reindex(['OIL','GAS'],fill_value=0).sum()/z.sum()),
            interpretation='Broad regional parent, not independently observed copper-mining technology.')
        before={(row['Item type'],row['DB Item'],row.get('DB Region')):float(row['Quantity']) for row in profiles[region]}
        for mode in ['nonnegative_block','all_factors']:
            candidate,details=parent_factor_proposal(profiles[region],factors,mode=mode)
            details.update(region=region,sector='MCU',parent=parent,output_MUSD=outputs[region],
                           coefficient_summary=summary(candidate))
            proposals.append(details)
            for row in candidate:
                key=row['Item type'],row['DB Item'],row.get('DB Region');old=before.get(key,0.)
                coefficient_rows.append(dict(region=region,proposal=mode,item_type=key[0],item=key[1],origin=key[2],
                    unit=row['Unit'],before=old,candidate=row['Quantity'],delta=float(row['Quantity'])-old))
    brazil=raw['BRA'];total=math.fsum(r['value'] for r in brazil.values())
    combined={'Capital':sum(brazil[k]['value'] for k in ['fixcap','othrent']),
              'Labor':sum(brazil[k]['value'] for k in ['lablskl','labmskl','labhskl']),
              'OTHTAX':brazil['othtax']['value'],'Tax':brazil['prodtax']['value']}
    bridge={k:dict(raw_value=q,raw_normalized=q/total,source_inventory=original_summary['factors'][k],
        source_to_raw_normalized_ratio=original_summary['factors'][k]/(q/total)) for k,q in combined.items()}
    if {str(f):sha(f) for f in files}!=hashes:raise RuntimeError('An input changed during the audit.')
    active=[r for r in regional if r['active_output']]
    low_v=[r for r in active if r['v_total']<.01]
    screening=dict(rule='Review flag only: net factor prior below 1% of output; not a certified anomaly or automatic replacement rule.',
        active_regions=len(active),zero_output_regions=len(regional)-len(active),
        regions_below_one_percent=low_v,
        share_of_mcu_output_in_flagged_regions=sum(r['output_MUSD'] for r in low_v)/sum(r['output_MUSD'] for r in active))
    donor_diagnostics={}
    for region in ['CHL','PER','AUS','USA']:
        values=raw.get(region,{})
        zsum=math.fsum(v['value'] for k,v in values.items() if k not in VA_CODES)
        total=math.fsum(v['value'] for v in values.values())
        if total and zsum:
            donor_diagnostics[region]=dict(raw_normalized_v_share=math.fsum(v['value'] for k,v in values.items() if k in VA_CODES)/total,
                raw_oil_gas_share_of_z=math.fsum(values.get(k,{'value':0})['value'] for k in ['oil','gas'])/zsum,
                interpretation='Source profile diagnostics only; not a selected or certified donor.')
    repo=Path(__file__).resolve().parents[1]
    report=dict(format='mcu-prior-review-v1',created_at=datetime.now(timezone.utc).isoformat(),status='review_and_proposals_only',
        source_sha256=hashes,code_revision=subprocess.check_output(['git','-C',str(repo),'rev-parse','HEAD'],text=True).strip(),
        code_sha256={str(f.relative_to(repo)):sha(f) for f in [Path(__file__).resolve(),repo/'src/entice_inventory/reconciliation/prior_review.py']},
        sources_unchanged=True,solver_invoked=False,preparation_policy_changed=False,
        raw_csv_units='Raw CSV values as supplied; normalized ratios are computed explicitly. No currency interpretation or conversion is asserted for raw numbers.',
        raw_region_count=len(raw),raw_invalid_regions=[r for r in native if r['invalid_positive_factor']],
        regions=regional,active_invalid_regions=active_bad,screening=screening,raw_same_activity_donor_diagnostics=donor_diagnostics,
        source_brazil_factor_rows={k:value for k,value in brazil.items() if k in VA_CODES},
        original_brazil=original_summary,resolved_brazil=resolved_summary,prepared_brazil=summary(profiles['BR']),
        raw_to_original_bridge=bridge,residual_source_rows=residual_rows,source_othtax_cluster_members=cluster_members,
        parent_references=parent_references,proposals=proposals,
        limitations=['Regional-parent V coefficients are donor assumptions, not newly measured child factor payments.',
            'Keeping the child Z composition also keeps its concentrated Oil/Gas structure. Repairing signs alone does not validate that technology.',
            'OTHTAX is supplied as an aggregate; the Tax/TTM split is an allocation, not two separate observations.',
            'Zero-output LT is flagged and excluded from active repair; do not use an invalid profile as a trusted donor.',
            'No add_sectors workbook, source, database or default preparation policy was changed; no solver was run.'])
    out.mkdir(parents=True)
    for name,values in [('raw-mcu-regions.csv',native),('prepared-mcu-regions.csv',regional),('candidate-coefficients.csv',coefficient_rows)]:
        with (out/name).open('w',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=list(values[0]));writer.writeheader();writer.writerows(values)
    (out/'review.json').write_text(json.dumps(report,indent=2,allow_nan=False))
    lines=['# MCU prior review','', '**Audit and proposals only. No model was run and no input workbook or preparation default was changed.**','',
        '## Provenance findings','',
        '- Negative Capital/Labor originate in the Purdue EXIOIOT CSV, before regional aggregation and before VA_fix.',
        '- MCU uses EXIOIOT; the COPCOSTS copper files belong to other copper activities.',
        '- Brazil Capital = FIXCAP + OTHRENT; Labor = LABLSKL + LABMSKL + LABHSKL. Every contributing source component is negative.',
        '- The original MARIO workbook multiplies normalized monetary priors by 0.99 and adds a 1% residual-sector cluster. This preserves the negative signs.',
        '- Capital/Labor are unchanged through the resolved D2.4 export. OTHTAX is split between Tax and TTM; PRODTAX is retained in Tax.',
        '- The raw source covers 122 MCU regions; Brazil and Lithuania have negative Capital/Labor. In the prepared 49-region geography, Brazil is the only affected region with positive MCU output (24 active, 25 zero-output regions).','',
        '## Brazil coefficient proposals','',
        '| Profile | Net V / output | Z / output | Oil+Gas / Z |','| --- | ---: | ---: | ---: |']
    rows=[('Current prepared prior',summary(profiles['BR']))]+[(p['mode'],p['coefficient_summary']) for p in proposals if p['region']=='BR']
    for label,s in rows:lines.append(f"| {label} | {100*s['v_total']:.6f}% | {100*s['z_total']:.6f}% | {100*s['oil_gas_share_of_z']:.3f}% |")
    ref=parent_references['BR'];lines.append(f"| Regional OXT reference | {100*ref['v_total']:.6f}% | {100*ref['z_total']:.6f}% | {100*ref['oil_gas_share_of_z']:.3f}% |")
    lines += ['', '`nonnegative_block` replaces Capital/Labor with regional OXT coefficients and retains the current child Tax/TTM. `all_factors` inherits all four regional OXT factor coefficients. Both rescale Z by one common factor, preserving its composition and origins, and close Z+V=1. Neither preserves the almost-zero source net V.','',
        'The proposals change only cost-coefficient profiles. Output and trade targets stay fixed; they are not readjusted to make a model solve. Capital/Labor remain nonnegative; Tax/TTM remain signed.','',
        '## Interpretation','',
        'These are reviewable imputation scenarios, not certified source corrections. In Brazil the source technology is concentrated in Oil/Gas; preserving that composition while repairing V does not independently validate the profile. The regional OXT technology is a useful control and a broad proxy, not a measured copper technology. Chile/Peru should not be chosen automatically as alternative donors merely because their factor signs are positive; donor technology and source coverage also require review.','']
    lines += ['', '## Broader screening','',
        f"{len(low_v)} of {len(active)} active MCU regions have net V priors below 1% of output. They account for {100*screening['share_of_mcu_output_in_flagged_regions']:.1f}% of prepared MCU output. This is a review threshold, not proof of an error or an automatic donor-replacement rule.",
        'The raw Chile and Peru profiles both have net V/output about 0.195% and Oil+Gas about 77.5% of Z before the 1% residual padding. A neighbouring-country donor does not by itself resolve the source-quality question. The regional OXT reference for Brazil has Oil+Gas about 0.523% of Z, versus 95.7% in the current child profile. Sector scope and valuation must be checked before treating either as a child observation.','']
    lines += ['- '+v for v in report['limitations']]
    lines += ['', '## Files','', '- `review.json`: hashes, lineage, raw source line numbers, factor values and full proposal metadata.', '- `raw-mcu-regions.csv`: 122 raw source profiles.', '- `prepared-mcu-regions.csv`: all 49 prepared regions and sign flags.', '- `candidate-coefficients.csv`: exact before/after coefficients for both Brazil proposals.', '', 'All source hashes were checked again after reading. Reproduce with `scripts/review_mcu_priors.py --prepared ORIGINAL_PREPARATION --output FRESH_DIRECTORY`.']
    (out/'README.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps(dict(output=str(out),active_invalid_regions=active_bad,parent_references=parent_references,proposals=proposals),indent=2,allow_nan=False))

if __name__=='__main__':main()
