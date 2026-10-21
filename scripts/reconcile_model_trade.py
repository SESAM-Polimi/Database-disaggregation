"""Project pending trade targets on signed row budgets of an accepted checkpoint.

Experimental copy-only preparation. The result requires a normal MARIO trial;
this projection alone is not a certificate of full MRIO feasibility.
"""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import shutil
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import numpy as np
import pandas as pd

from run_model_trial import fingerprints
from entice_inventory.reconciliation.row_trade import fit_row_trade

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def rewrite_sheets(source, destination, frames):
    """Preserve every ZIP member outside the explicitly replaced sheets."""
    from openpyxl.utils import get_column_letter
    with ZipFile(source) as before:
        rels = {e.attrib['Id']:e.attrib['Target'] for e in ET.fromstring(before.read('xl/_rels/workbook.xml.rels'))}
        names = {}
        for e in ET.fromstring(before.read('xl/workbook.xml')).find(f'{{{NS}}}sheets'):
            target = rels[e.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            names[e.attrib['name']] = target.lstrip('/') if target.startswith('/') else 'xl/'+target
        replacements = {}
        for name,frame in frames.items():
            root = ET.fromstring(before.read(names[name]))
            data = root.find(f'{{{NS}}}sheetData')
            data.clear()
            for i,values in enumerate([frame.columns.tolist(),*frame.itertuples(index=False,name=None)],1):
                row = ET.SubElement(data,f'{{{NS}}}row',r=str(i))
                for j,value in enumerate(values,1):
                    if pd.isna(value):
                        continue
                    cell = ET.SubElement(row,f'{{{NS}}}c',r=f'{get_column_letter(j)}{i}')
                    if isinstance(value,(int,float,np.number)):
                        ET.SubElement(cell,f'{{{NS}}}v').text = repr(float(value))
                    else:
                        cell.set('t','inlineStr')
                        ET.SubElement(ET.SubElement(cell,f'{{{NS}}}is'),f'{{{NS}}}t').text = str(value)
            dimension = root.find(f'{{{NS}}}dimension')
            extent = f'A1:{get_column_letter(len(frame.columns))}{len(frame)+1}'
            if dimension is not None:
                dimension.set('ref',extent)
            autofilter = root.find(f'{{{NS}}}autoFilter')
            if autofilter is not None:
                autofilter.set('ref',extent)
            replacements[names[name]] = ET.tostring(root,encoding='utf-8',xml_declaration=True)
        with ZipFile(destination,'w') as after:
            for info in before.infolist():
                after.writestr(copy.copy(info),replacements.get(info.filename,before.read(info.filename)))
    with ZipFile(source) as before, ZipFile(destination) as after:
        assert set(before.namelist()) == set(after.namelist())
        assert all(before.read(n)==after.read(n) for n in before.namelist() if n not in replacements)
    return sorted(replacements)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--checkpoint',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--trade-audit',type=Path,required=True)
    parser.add_argument('--policy',type=Path,required=True)
    args = parser.parse_args()
    source, output = args.checkpoint.resolve(),args.output.resolve()
    if output==source or source in output.parents:
        parser.error('Output must be outside the immutable checkpoint.')
    output.mkdir(parents=True,exist_ok=False)
    hashes = fingerprints(source)
    from mario.ops.add_sector_steps import load_checkpoint
    db = load_checkpoint(source)
    state = db._intermediate_steps_state
    if state['split_model_options']['accounting'] != 'inherit':
        raise ValueError('This projection requires inherited accounting.')
    workbook = source/'remaining_add_sectors.xlsx'
    master, outputs, trades = (pd.read_excel(workbook,sheet_name=s) for s in ('Master','Total outputs','Trades'))
    parents = master.set_index('Sector')['Parent Sector'].to_dict()
    if set(parents)&set(state['completed_sectors']):
        raise ValueError('Expected only pending sectors in the remaining workbook.')
    audit = pd.read_csv(args.trade_audit).set_index(['sector','origin','destination'])
    policy = json.loads(args.policy.read_text())
    source_weights = policy['trade']['weights']
    q = outputs.set_index(['Sector','Region']).Quantity
    old = trades.set_index(['Sector_from','Region_from','Region_to'])
    regions = db.get_index('Region')
    records, changes, added = [],[],[]
    for parent in sorted(set(parents.values())):
        children = sorted(s for s in parents if parents[s]==parent)
        for origin in regions:
            positive,negative = np.zeros(len(regions)),np.zeros(len(regions))
            for table in (db.Z,db.Y):
                row = table.loc[(origin,'Sector',parent)]
                positive += row.clip(lower=0).groupby(level='Region').sum().reindex(regions).to_numpy()
                negative += (-row.clip(upper=0)).groupby(level='Region').sum().reindex(regions).to_numpy()
            parent_x = float(db.X.loc[(origin,'Sector',parent)].iloc[0])
            child_x = np.array([q.loc[(c,origin)] for c in children])
            residual_x = parent_x-child_x.sum()
            bound = (len(children)+1)*np.finfo(float).eps*(abs(parent_x)+abs(child_x.sum()))
            if residual_x < -bound:
                raise ValueError(f'Output excess for {parent}/{origin}.')
            xs = np.r_[max(residual_x,0.),child_x]
            discrepancy = (positive-negative).sum()-parent_x
            row_targets = xs*(1+discrepancy/parent_x) if parent_x>0 else xs.copy()
            prior, weights, observed = np.zeros((len(children),len(regions))),np.full((len(children),len(regions)),.1),np.zeros((len(children),len(regions)),bool)
            for i,child in enumerate(children):
                for j,destination in enumerate(regions):
                    if origin==destination:
                        continue
                    key = child,origin,destination
                    if key in old.index:
                        prior[i,j] = old.loc[key,'Quantity']
                        observed[i,j] = True
                        if key not in audit.index or not np.isclose(audit.loc[key,'reconciled'],prior[i,j],atol=1e-6,rtol=1e-8):
                            raise ValueError(f'Trade source audit does not match workbook: {key}.')
                        weights[i,j] = source_weights[audit.loc[key,'source_kind']]
            print(f'Checking {parent}/{origin}',flush=True)
            try:
                fitted, record = fit_row_trade(positive,negative,row_targets,xs,prior,weights,regions.index(origin),observed=observed)
            except Exception:
                np.savez(output/'failed_group.npz',positive=positive,negative=negative,row_targets=row_targets,
                         outputs=xs,priors=prior,weights=weights,origin=regions.index(origin),observed=observed)
                raise
            record.update(parent=parent,origin=origin)
            records.append(record)
            if record['status']=='already_row_feasible':
                continue
            for i,child in enumerate(children):
                for j,destination in enumerate(regions):
                    if destination==origin:
                        continue
                    key = child,origin,destination
                    value = float(fitted[i,j])
                    kind = audit.loc[key,'source_kind'] if key in audit.index else 'imputed'
                    changes.append(dict(sector=child,parent=parent,origin=origin,destination=destination,
                                        prior=float(prior[i,j]),reconciled=value,source_kind=kind,previously_observed=bool(observed[i,j])))
                    if key in old.index:
                        old.loc[key,'Quantity'] = value
                        old.loc[key,'Source'] = 'signed-row-trade-projection-v1'
                    else:
                        added.append(dict(Sector_from=child,Region_from=origin,Region_to=destination,Quantity=value,
                                          Unit=db.units['Sector'].loc[parent,'unit'],Source='signed-row-trade-projection-v1',
                                          Notes='Fitted route; prior absent. See reconciliation audit.'))
    result = pd.concat([old.reset_index(),pd.DataFrame(added)],ignore_index=True).reindex(columns=trades.columns)
    plan = pd.DataFrame({'Reconciled remaining sectors':list(parents)})
    destination = output/'add_sectors.xlsx'
    changed_members = rewrite_sheets(workbook,destination,{'Trades':result,'Intermediate steps':plan})
    assert fingerprints(source)==hashes
    pd.DataFrame(changes).to_csv(output/'trade_changes.csv',index=False)
    report = dict(status='projected',full_mrio_certified=False,source_checkpoint=str(source),
                  source_checkpoint_sha256=hashes,source_checkpoint_unchanged=True,
                  source_trade_audit=str(args.trade_audit.resolve()),source_policy=str(args.policy.resolve()),
                  source_trade_audit_sha256=hashlib.sha256(args.trade_audit.read_bytes()).hexdigest(),
                  source_policy_sha256=hashlib.sha256(args.policy.read_bytes()).hexdigest(),
                  workbook_sha256=hashlib.sha256(destination.read_bytes()).hexdigest(),
                  only_changed_zip_members=changed_members,all_other_workbook_members_byte_identical=True,
                  outputs_unchanged=True,source_weights=source_weights,unobserved_route_weight=.1,
                  prior_floor='max(abs(prior), 0.01*parent_route_capacity, 1e-6 MUSD)',
                  permits_prior_increases=True,groups=records)
    (output/'reconciliation.json').write_text(json.dumps(report,indent=2,allow_nan=False))
    shutil.copy2(__file__,output/'driver.py')
    print(json.dumps({k:report[k] for k in ('status','full_mrio_certified','outputs_unchanged','source_checkpoint_unchanged')},indent=2))


if __name__=='__main__':
    main()
